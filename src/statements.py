"""MISB PDF statements: exact decimal ledger amounts and reference artwork."""
import io
from zipfile import BadZipFile
from datetime import datetime, date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import openpyxl

CENT = Decimal('0.01')
SERVICE_RATE = Decimal('0.20')
SST_RATE = Decimal('0.08')
SST_START = date(2026, 8, 1)
MONTHS = ('January','February','March','April','May','June','July','August','September','October','November','December')


def money(value):
    try:
        amount = Decimal(str(value))
        if not amount.is_finite():
            raise ValueError('Amounts must be finite')
        return amount.quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError):
        raise ValueError('A transaction has an invalid amount') from None


def timestamp(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    for fmt in ('%d-%b-%Y %H:%M:%S','%Y-%m-%d %H:%M:%S','%Y-%m-%d','%d/%m/%Y'):
        try:
            return datetime.strptime(str(value).strip(),fmt)
        except ValueError:
            pass
    raise ValueError('A successful transaction has an invalid date')


def profit_components(net_amount, transaction_date):
    """Split a net profit payout using the SST rules effective on its transaction date."""
    net = money(net_amount)
    taxable = transaction_date >= SST_START
    divisor = Decimal('0.784') if taxable else Decimal('0.80')
    gross = (net / divisor).quantize(CENT, rounding=ROUND_HALF_UP)
    service = (gross * SERVICE_RATE).quantize(CENT, rounding=ROUND_HALF_UP)
    sst = (service * SST_RATE).quantize(CENT, rounding=ROUND_HALF_UP) if taxable else Decimal(0)
    return gross, service, sst


def statement_data(data, as_of=None):
    wb = openpyxl.load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    transactions = []
    try:
        ws = wb.worksheets[0]
        headers = next(ws.iter_rows(min_row=2,max_row=2,values_only=True))
        for index, values in enumerate(ws.iter_rows(min_row=3,values_only=True)):
            row = dict(zip(headers,values))
            if str(row.get('Status','')).upper() != 'SUCCESSFUL':
                continue
            dt = timestamp(row.get('Date'))
            note = row.get('Note #')
            note = '-' if note in (None,'',0,'0') else str(note).removeprefix('IIF-')
            if note.endswith('.0'): note=note[:-2]
            transactions.append({'timestamp':dt,'order':int(row.get('ID') or index),
                'description':str(row.get('Action') or '').strip(),'note':note,
                'previous':money(row.get('Previous Balance(MYR)')),
                'amount':money(row.get('Sum Involved(MYR)')),
                'current':money(row.get('Current Balance(MYR)'))})
            if len(transactions)>10000:
                raise ValueError('Statements support up to 10,000 transactions per workbook')
    finally:
        wb.close()
    if not transactions:
        raise ValueError('The workbook has no successful transactions')
    transactions.sort(key=lambda r:(r['timestamp'],r['order']))
    try:
        cutoff = date.fromisoformat(as_of) if as_of else transactions[-1]['timestamp'].date()
    except (ValueError,TypeError):
        raise ValueError('Cut-off date must be YYYY-MM-DD') from None
    selected = [r for r in transactions if r['timestamp'].date() <= cutoff]
    rows, warnings = [], []
    for r in selected:
        r = dict(r)
        action = r['description']
        if action == 'Deposit':
            continue
        if action == 'Withdrawal':
            # The export records the RM1 fee as a separate Withdrawal row.
            # Match its balance movement to the approved withdrawal, not a request.
            fee = (r['amount'] == Decimal('1.00')
                   and r['previous'] - r['current'] == r['amount']
                   and any(a['description'] == 'Withdrawal Approved'
                           and a['timestamp'] == r['timestamp']
                           and a['current'] == r['previous'] for a in selected))
            if not fee:
                continue
            r['description'] = 'Withdrawal Fee'
        if action in ('Deposit Approved', 'Withdrawal Approved'):
            r['description'] = action.removesuffix(' Approved')
            if r['note'] == '-':
                r['note'] = '0'
        rows.append(r)
    if not rows:
        raise ValueError('No completed cash transactions exist on or before this cut-off date')
    breaks = sum(a['current'] != b['previous'] for a,b in zip(rows,rows[1:]))
    if breaks:
        warnings.append(f'{breaks} balance discontinuity/discontinuities found. The PDF preserves the original ledger balances; check that the log is complete.')
    totals = {
        'starting_balance':rows[0]['previous'],'ending_balance':rows[-1]['current'],
        'total_investment':sum((r['amount'] for r in rows if r['description']=='Investment Committed'),Decimal(0)),
        'principal_received':sum((r['amount'] for r in rows if r['description']=='Principal Payout'),Decimal(0)),
        'nett_returns':sum((r['amount'] for r in rows if r['description']=='Profit Payout'),Decimal(0))}
    detailed, gross_total, service_total, sst_total = [], Decimal(0), Decimal(0), Decimal(0)
    for r in rows:
        if r['description'] != 'Profit Payout':
            detailed.append(r)
            continue
        taxable = r['timestamp'].date() >= SST_START
        gross, service, sst = profit_components(r['amount'], r['timestamp'].date())
        common = {'timestamp':r['timestamp'],'order':r['order'],'note':r['note']}
        detailed.append({**common,'description':'Gross Profit','previous':r['previous'],
                         'amount':gross,'current':None})
        detailed.append({**common,'description':'Service Charge','previous':None,
                         'amount':-service,'current':None})
        if taxable:
            detailed.append({**common,'description':'SST','previous':None,
                             'amount':-sst,'current':None})
        detailed.append({**common,'description':'Net Profit','previous':None,
                         'amount':r['amount'],'current':r['current']})
        gross_total += gross; service_total += service; sst_total += sst
    totals.update(total_gross_returns=gross_total, service_fee=service_total, sst=sst_total)
    def value(amount):
        return '' if amount is None else format(amount,'.2f')
    return {'as_of':cutoff.isoformat(),'latest_date':transactions[-1]['timestamp'].date().isoformat(),
        'investor_id':'5490','investor_name':'Amanahraya Trustees Berhad',
        'summary':{k:format(v,'.2f') for k,v in totals.items()},
        'rows':[{'date':r['timestamp'].strftime('%d-%b-%Y %H:%M:%S'),'description':r['description'],
                 'note':r['note'],**{k:value(r[k]) for k in ('previous','amount','current')}} for r in detailed],
        'row_count':len(detailed),'page_count':1+max(0,(len(detailed)-33+57)//58),'warnings':warnings}


def statement_pdf(statement):
    # Lazy imports keep the Worker startup path small.
    from pypdf import PdfReader, PdfWriter
    from pypdf._cmap import get_encoding
    from pypdf.generic import NameObject, DecodedStreamObject
    from statement_template import TEMPLATE
    reader=PdfReader(io.BytesIO(TEMPLATE)); writer=PdfWriter()
    first=writer.add_page(reader.pages[0])
    fonts=first['/Resources']['/Font']

    def font_info(name):
        font=fonts[name].get_object()
        if '/DescendantFonts' not in font:
            return None
        reverse={value:key for key,value in get_encoding(font)[1].items()}
        descendant=font['/DescendantFonts'][0].get_object(); widths={}
        values=descendant.get('/W',[])
        if hasattr(values,'get_object'): values=values.get_object()
        i=0
        while i < len(values):
            start=int(values[i]); following=values[i+1]
            if isinstance(following,list):
                for offset,item in enumerate(following): widths[start+offset]=float(item)
                i+=2
            else:
                end=int(following); item=float(values[i+2])
                for code in range(start,end+1): widths[code]=item
                i+=3
        return reverse,widths,float(descendant.get('/DW',1000))

    font_data={name:font_info(name) for name in ('/F1','/F3')}

    def encoded_text(value,font):
        info=font_data.get(font)
        if info is None:
            raw=value.encode('cp1252')
            f=fonts[font].get_object(); widths=f['/Widths']; start=int(f['/FirstChar'])
            return raw,sum(float(widths[c-start]) if 0<=c-start<len(widths) else 0 for c in raw)
        reverse,widths,default=info; raw=bytearray(); total=0
        for character in value:
            code_text=reverse.get(character)
            if code_text is None:
                raise ValueError(f'The statement font cannot render {character!r}')
            for item in code_text:
                code=ord(item);raw.extend(code.to_bytes(2,'big'));total+=widths.get(code,default)
        return bytes(raw),total

    def width(value,font,size):
        return encoded_text(value,font)[1]*size/1000

    def text(text,x,y,size=6.96,font='/F1',align='left'):
        w=width(text,font,size)
        if align=='right': x-=w
        if align=='center': x-=w/2
        literal=encoded_text(text,font)[0].hex()
        return f'BT 0 g {font} {size} Tf 1 0 0 1 {x:.4f} {792-y:.4f} Tm <{literal}> Tj ET\n'

    cutoff=date.fromisoformat(statement['as_of'])
    title=f'ACCOUNT STATEMENT (YEAR TO DAY- {cutoff.day} {MONTHS[cutoff.month-1]} {cutoff:%y})'
    first_ops=text(title,302.5,190.8,11.4,'/T2','center')
    for key,y in zip(('starting_balance','ending_balance','total_investment','principal_received','total_gross_returns'),(261.96,273.60,285.24,296.88,308.52)):
        first_ops+=text(format(Decimal(statement['summary'][key]),',.2f'),312.95,y,7.56,font='/F1',align='right')
    for key,y in zip(('service_fee','sst','nett_returns'),(261.96,273.60,285.24)):
        first_ops+=text(format(Decimal(statement['summary'][key]),',.2f'),550.33,y,7.56,font='/F1',align='right')
    all_rows=statement['rows']; groups=[all_rows[:33]]+[all_rows[i:i+58] for i in range(33,len(all_rows),58)]
    columns=(50.964,170.65,315.85,360.73,418.45,487.21,553.5)

    def display_amount(value):
        if value == '': return ''
        amount=Decimal(value)
        return f'({abs(amount):,.2f})' if amount < 0 else f'{amount:,.2f}'
    for page_index,rows in enumerate(groups):
        page=first if page_index==0 else writer.add_blank_page(width=612,height=792)
        if page_index: page[NameObject('/Resources')]=first['/Resources']
        ops=first_ops if page_index==0 else ''
        top=353.64 if page_index==0 else 54.24
        bottom=top+len(rows)*11.64
        ops+='0 G 0.6 w\n'
        for x in columns:
            ops+=f'{x} {792-top:.4f} m {x} {792-bottom:.4f} l S\n'
        for i in range(len(rows)+1):
            y=792-top-i*11.64
            ops+=f'{columns[0]} {y:.4f} m {columns[-1]} {y:.4f} l S\n'
        for i,row in enumerate(rows):
            y=(360.60 if page_index==0 else 62.52)+11.64*i
            for value,left,right in zip((row['date'],row['description'],row['note']),columns[:3],columns[1:4]):
                size=min(6.96,6.96*(right-left-3)/max(1,width(value,'/F1',6.96)))
                ops+=text(value,(left+right)/2,y,size,font='/F1',align='center')
            for key,right in zip(('previous','amount','current'),(409.35,486.85,550.40)):
                value=display_amount(row[key]); font='/F3' if value.startswith('(') else '/F1'
                ops+=text(value,right,y,font=font,align='right')
        old=page.get_contents()
        stream=DecodedStreamObject()
        stream.set_data((b'q\n'+old.get_data()+b'\nQ\n' if old is not None else b'')+ops.encode('ascii'))
        page[NameObject('/Contents')]=stream
        page.compress_content_streams()
    writer.add_metadata({'/Title':f'MISB Account Statement through {statement["as_of"]}', '/Author':'Crowd Sense Sdn Bhd'})
    output=io.BytesIO();writer.write(output);output.seek(0)
    return output


def register_statement_routes(app):
    from flask import request, jsonify, send_file, g
    from storage import (read_source, source_rows, save_source, source_exists,
                         save_parsed_source, audit_event)
    from uploads import inspect_workbook

    def build():
        if not source_exists('transactions'):
            raise ValueError('Upload a transaction log first')
        expected=request.args.get('version')
        if expected and expected!=source_rows()['transactions']['object_key']:
            return None
        result=statement_data(read_source('transactions').read(),request.args.get('as_of'))
        result['version']=source_rows()['transactions']['object_key']
        return result

    @app.route('/api/account-statement',methods=['GET','POST'])
    def preview_statement():
        try:
            if request.method=='POST':
                file=request.files.get('file')
                if not file or not file.filename.lower().endswith('.xlsx'):
                    raise ValueError('Choose a transaction log in .xlsx format')
                data=file.read()
                from app import load_transactions
                loaded=load_transactions() if source_exists('transactions') else []
                current=list(getattr(g,'transactions_base',loaded))
                quality=inspect_workbook('transactions',data,current)
                if not quality.get('can_import'):
                    return jsonify(error='The transaction workbook failed data-quality checks.',
                                   issues=quality.get('issues',[])),422
                if (quality.get('comparison') or {}).get('requires_acknowledgement'):
                    return jsonify(error='This ledger is older or removes live transactions. Review and import it from Data Sources if that rollback is intentional.',
                                   comparison=quality['comparison']),409
                result=statement_data(data,request.form.get('as_of'))
                save_source('transactions',data,result['latest_date'],file.filename)
                g.pop('transactions_base',None);g.pop('transactions_effective',None)
                load_transactions()
                save_parsed_source('transactions',source_rows()['transactions']['object_key'],g.transactions_base)
                audit_event('Imported statement ledger','source','transactions',{'as_of':result['latest_date']})
                result['version']=source_rows()['transactions']['object_key']
            else:
                result=build()
            if result is None: return jsonify(error='The ledger has changed. Refresh the statement preview.'),409
            return jsonify(result)
        except (BadZipFile,openpyxl.utils.exceptions.InvalidFileException):
            return jsonify(error='This file is not a readable Excel workbook.'),400
        except (ValueError,KeyError,StopIteration) as error:
            return jsonify(error=str(error)),400

    @app.get('/api/export/account-statement.pdf')
    def export_statement_pdf():
        try:
            result=build()
            if result is None: return jsonify(error='The ledger has changed. Refresh the statement preview.'),409
            filename=f'MISB_Account_Statement_{result["as_of"]}.pdf';pdf=statement_pdf(result)
            from storage import save_statement_report_run
            summary=result.get('summary',{})
            save_statement_report_run(result['as_of'],result['version'],filename,{
                'row_count':result.get('row_count',0),'page_count':result.get('page_count',0),
                'opening_balance':summary.get('starting_balance',0),'closing_balance':summary.get('ending_balance',0),
                'gross_returns':summary.get('total_gross_returns',0)})
            audit_event('Generated account statement','report',result['as_of'],{
                'filename':filename,'source_version':result['version'],'rows':result.get('row_count',0)})
            return send_file(pdf,mimetype='application/pdf',as_attachment=True,download_name=filename)
        except (ValueError,KeyError) as error:
            return jsonify(error=str(error)),400

    @app.get('/api/account-statement-runs')
    def account_statement_history():
        from storage import statement_report_runs
        return jsonify(statement_report_runs(request.args.get('limit',20)))

    @app.get('/api/account-statement-runs/<int:run_id>/pdf')
    def reproduce_account_statement(run_id):
        from storage import statement_report_run,read_source_version
        run=statement_report_run(run_id)
        if not run: return jsonify(error='That statement history record was not found.'),404
        try:
            source,_=read_source_version('transactions',run['source_object_key'])
            result=statement_data(source.read(),run['as_of'])
            audit_event('Recreated account statement','report',run['as_of'],{
                'history_id':run_id,'filename':run['filename'],'source_version':run['source_object_key']})
            return send_file(statement_pdf(result),mimetype='application/pdf',as_attachment=True,
                             download_name=run['filename'])
        except (ValueError,KeyError) as error:
            return jsonify(error=str(error)),400
