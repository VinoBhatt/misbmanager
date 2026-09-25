"""Bound spreadsheet expansion before handing XML to openpyxl."""
import io
import zipfile
from datetime import date, datetime
import openpyxl
import re


def validate_workbook(kind, data):
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Workbook exceeds 10 MB')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
            raise ValueError('Expanded workbook is too large')
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if kind == 'transactions':
            sheet = wb.worksheets[0]
            headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
            required = {'ID','Action','Sum Involved(MYR)','Previous Balance(MYR)',
                        'Current Balance(MYR)','Note #','Status','Date'}
        elif kind == 'simulation':
            if 'Query result' not in wb.sheetnames:
                raise ValueError('Missing Query result sheet')
            sheet = wb['Query result']
            headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            required = {'Loan Code','Issuer Name','Investment Amount','Loan Status',
                        'Paid Principal','Unpaid Principal','Paid Profit','Unpaid Profit'}
        else:
            if not {'Sheet1','Sheet2'}.issubset(wb.sheetnames):
                raise ValueError('Missing MIDAS sheets')
            return
        if not required.issubset(set(headers)):
            raise ValueError('Required columns are missing')
        if kind == 'simulation':
            legacy={'Gross Profit','Late Profit','Service Fee ','Net Profit'}
            expanded={'Gross Profit Earned','Total Gross Profit','Late Payment Charges','Service Fee ','SST','Net Profit','Installment'}
            if not (legacy.issubset(set(headers)) or expanded.issubset(set(headers))):
                raise ValueError('The simulation profit columns are not a supported format')
        if (sheet.max_row or 0) > 50000 or (sheet.max_column or 0) > 500:
            raise ValueError('Workbook dimensions exceed supported limits')
    finally:
        wb.close()


def inspect_workbook(kind, data, current=None):
    """Return a concise, non-mutating quality report for an import candidate."""
    validate_workbook(kind,data)
    wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    issues=[];stats={'bytes':len(data),'sheets':len(wb.sheetnames)}
    def issue(level,message,count=1):
        issues.append({'level':level,'message':message,'count':count})
    def parsed_date(value):
        if isinstance(value,datetime): return value
        if isinstance(value,date): return datetime.combine(value,datetime.min.time())
        for fmt in ('%d-%b-%Y %H:%M:%S','%d/%m/%Y','%Y-%m-%d %H:%M:%S','%Y-%m-%d'):
            try: return datetime.strptime(str(value or '').strip(),fmt)
            except ValueError: pass
        return None
    def note_code(value):
        if value is None: return ''
        if isinstance(value,(int,float)) and not isinstance(value,bool): return f'IIF-{int(value)}'
        text=str(value).replace('**','').strip().upper();match=re.search(r'(?:IIF[-\s]*)?(\d{3,6})',text)
        return f'IIF-{match.group(1)}' if match else text
    def number(value):
        try: return float(value or 0)
        except (TypeError,ValueError): return 0.0
    def changed(before,after,fields):
        for field in fields:
            left=before.get(field);right=after.get(field)
            if isinstance(left,(int,float)) or isinstance(right,(int,float)):
                try:
                    if abs(float(left or 0)-float(right or 0))>.005: return True
                except (TypeError,ValueError): return True
            elif str(left or '').strip()!=str(right or '').strip(): return True
        return False
    def comparison(before,after,fields,label):
        previous={str(row.get('id') or row.get('loan_code') or ''):row for row in (before or [])
                  if row.get('id') not in (None,'') or row.get('loan_code')}
        candidate={str(row.get('id') or row.get('loan_code') or ''):row for row in after
                   if row.get('id') not in (None,'') or row.get('loan_code')}
        added=sorted(set(candidate)-set(previous));removed=sorted(set(previous)-set(candidate))
        altered=sorted(key for key in set(previous)&set(candidate) if changed(previous[key],candidate[key],fields))
        return {'entity':label,'current_rows':len(previous),'candidate_rows':len(candidate),
                'added':len(added),'removed':len(removed),'changed':len(altered),
                'added_sample':added[:8],'removed_sample':removed[:8],'changed_sample':altered[:8]}
    try:
        if kind=='transactions':
            ws=wb.worksheets[0];headers=next(ws.iter_rows(min_row=2,max_row=2,values_only=True))
            rows=[]
            for values in ws.iter_rows(min_row=3,values_only=True):
                if any(value is not None for value in values): rows.append(dict(zip(headers,values)))
            successful=[row for row in rows if str(row.get('Status','')).upper()=='SUCCESSFUL']
            stats.update(rows=len(rows),successful=len(successful),pending=len(rows)-len(successful))
            ids=[row.get('ID') for row in rows if row.get('ID') not in (None,'')]
            duplicates=len(ids)-len(set(ids))
            if duplicates: issue('error','Duplicate transaction IDs',duplicates)
            missing_ids=sum(row.get('ID') in (None,'') for row in successful)
            if missing_ids: issue('error','Successful transactions without an ID',missing_ids)
            invalid_dates=sum(parsed_date(row.get('Date')) is None for row in successful)
            if invalid_dates: issue('error','Successful transactions with invalid dates',invalid_dates)
            invalid_amounts=sum(not isinstance(row.get('Sum Involved(MYR)'),(int,float)) for row in successful)
            if invalid_amounts: issue('error','Successful transactions with invalid amounts',invalid_amounts)
            balance_columns=('Previous Balance(MYR)','Current Balance(MYR)')
            invalid_balances=sum(any(not isinstance(row.get(column),(int,float)) for column in balance_columns) for row in successful)
            if invalid_balances: issue('error','Successful transactions with invalid balances',invalid_balances)
            note_actions={'Investment Committed','Principal Payout','Profit Payout'}
            missing_notes=sum(row.get('Action') in note_actions and row.get('Note #') in (None,'',0,'0') for row in successful)
            if missing_notes: issue('warning','Note transactions without a note number',missing_notes)
            approved=sum(row.get('Action') in ('Deposit Approved','Withdrawal Approved') for row in successful)
            # The platform emits zero-balance Deposit/Withdrawal shadow entries beside
            # the Approved entry. They remain visible for traceability but never drive cash.
            shadow_actions={'Deposit','Withdrawal'}
            effective=[row for row in successful if row.get('Action') not in shadow_actions]
            shadows=len(successful)-len(effective)
            stats.update(approved_cash_movements=approved,excluded_cash_shadows=shadows)
            valid_ledger=[row for row in successful if parsed_date(row.get('Date')) is not None and
                          all(isinstance(row.get(column),(int,float)) for column in balance_columns)]
            def id_number(row):
                try: return float(row.get('ID') if row.get('ID') not in (None,'') else row.get('id'))
                except (TypeError,ValueError): return 0
            ordered=sorted(valid_ledger,key=lambda row:(parsed_date(row.get('Date')),id_number(row)),reverse=True)
            breaks=sum(abs(float(newer['Previous Balance(MYR)'])-float(older['Current Balance(MYR)']))>.01
                       for newer,older in zip(ordered,ordered[1:]))
            stats['balance_breaks']=breaks
            if breaks: issue('error','Ledger balance is not continuous between transactions',breaks)
            movement_rows=[row for row in effective if isinstance(row.get('Sum Involved(MYR)'),(int,float)) and
                           all(isinstance(row.get(column),(int,float)) for column in balance_columns)]
            mismatches=sum(abs(abs(float(row['Current Balance(MYR)'])-float(row['Previous Balance(MYR)']))-
                               abs(float(row['Sum Involved(MYR)'])))>.01 for row in movement_rows)
            stats['movement_mismatches']=mismatches
            if mismatches: issue('error','Transaction amount does not match its balance movement',mismatches)
            fingerprints=[]
            for row in effective:
                when=parsed_date(row.get('Date'))
                fingerprints.append((when.isoformat() if when else '',str(row.get('Action') or '').strip(),
                                     float(row.get('Sum Involved(MYR)') or 0),str(row.get('Note #') or '').strip(),
                                     str(row.get('Transaction No') or '').strip(),str(row.get('Pay Via') or '').strip()))
            repeated=len(fingerprints)-len(set(fingerprints))
            stats['repeated_movements']=repeated
            if repeated: issue('error','Repeated financial movements with different transaction IDs',repeated)
            candidate_rows=[{'id':row.get('ID'),'action':row.get('Action') or '',
                             'amount':number(row.get('Sum Involved(MYR)')),
                             'previous_balance':number(row.get('Previous Balance(MYR)')),
                             'current_balance':number(row.get('Current Balance(MYR)')),
                             'note':note_code(row.get('Note #')),'status':row.get('Status') or '',
                             'date':parsed_date(row.get('Date')).strftime('%Y-%m-%d') if parsed_date(row.get('Date')) else ''}
                            for row in rows]
            source_comparison=comparison(current,candidate_rows,
                ('action','amount','previous_balance','current_balance','note','status','date'),'transactions') if current is not None else None
            if source_comparison is not None:
                candidate_latest=max(candidate_rows,key=lambda row:(row['date'],id_number(row)),default={})
                current_latest=max(current or [],key=lambda row:(str(row.get('date') or ''),id_number(row)),default={})
                source_comparison.update({'current_latest_date':current_latest.get('date') or '',
                    'candidate_latest_date':candidate_latest.get('date') or '',
                    'current_balance':number(current_latest.get('current_balance')),
                    'candidate_balance':number(candidate_latest.get('current_balance'))})
                reasons=[]
                if source_comparison['removed']:
                    reasons.append(f"The candidate omits {source_comparison['removed']} existing transaction(s).")
                if (source_comparison['current_latest_date'] and source_comparison['candidate_latest_date'] and
                        source_comparison['candidate_latest_date']<source_comparison['current_latest_date']):
                    reasons.append('The candidate ledger ends before the current live ledger.')
                source_comparison.update(requires_acknowledgement=bool(reasons),regression_reasons=reasons)
        elif kind=='simulation':
            ws=wb['Query result'];headers=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
            rows=[dict(zip(headers,values)) for values in ws.iter_rows(min_row=2,values_only=True)
                  if values and values[0] is not None and values[1]]
            codes=[str(row.get('Loan Code') or '').strip().upper() for row in rows]
            duplicates=len(codes)-len(set(codes))
            missing_issuers=sum(not str(row.get('Issuer Name') or '').strip() for row in rows)
            invalid_amounts=sum(not isinstance(row.get('Investment Amount'),(int,float)) or row.get('Investment Amount',0)<0 for row in rows)
            stats.update(rows=len(rows),format=('Expanded profit' if 'Total Gross Profit' in headers else 'Original'))
            if duplicates: issue('error','Duplicate loan codes',duplicates)
            if missing_issuers: issue('warning','Notes without an issuer name',missing_issuers)
            if invalid_amounts: issue('error','Notes with invalid investment amounts',invalid_amounts)
            candidate_rows=[{'loan_code':note_code(row.get('Loan Code')),
                             'issuer':row.get('Issuer Name') or '',
                             'investment_amount':number(row.get('Investment Amount')),
                             'loan_status':row.get('Loan Status') or '',
                             'paid_principal':number(row.get('Paid Principal')),
                             'unpaid_principal':number(row.get('Unpaid Principal')),
                             'paid_profit':number(row.get('Paid Profit')),
                             'unpaid_profit':number(row.get('Unpaid Profit'))} for row in rows]
            source_comparison=comparison(current,candidate_rows,('issuer','investment_amount','loan_status',
                'paid_principal','unpaid_principal','paid_profit','unpaid_profit'),'notes') if current is not None else None
            if source_comparison is not None:
                reasons=[]
                if source_comparison['removed']:
                    reasons.append(f"The candidate omits {source_comparison['removed']} existing simulation note(s).")
                source_comparison.update(requires_acknowledgement=bool(reasons),regression_reasons=reasons)
        else:
            stats.update(rows=(wb['Sheet2'].max_row or 0),format='MIDAS projections')
            source_comparison=None
        if not issues: issue('ok','No blocking data-quality issues found',0)
        result={'kind':kind,'stats':stats,'issues':issues,
                'can_import':not any(item['level']=='error' for item in issues)}
        if source_comparison is not None: result['comparison']=source_comparison
        return result
    finally:
        wb.close()
