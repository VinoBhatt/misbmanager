from flask import Flask, jsonify, request, send_file, g
from pathlib import Path
from copy import copy
from datetime import datetime, date, timedelta
import openpyxl, json, re, io, csv, calendar
from storage import (connect, read_source, source_exists, source_rows, save_source,
                     read_parsed_source, save_parsed_source, source_versions,
                     select_source_version, audit_event, audit_rows)

app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024


def num(v):
    try: return float(v or 0)
    except: return 0.0

def iso(v):
    if isinstance(v, (datetime, date)): return v.strftime('%Y-%m-%d')
    if v is None: return None
    s=str(v).strip()
    for fmt in ('%d-%b-%Y %H:%M:%S','%d/%m/%Y','%Y-%m-%d %H:%M:%S','%Y-%m-%d'):
        try: return datetime.strptime(s,fmt).strftime('%Y-%m-%d')
        except: pass
    return s

def normalize_note_code(v):
    if v is None: return ''
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try: return f'IIF-{int(v)}'
        except: pass
    s=str(v).replace('**','').strip().upper()
    m=re.search(r'(?:IIF[-\s]*)?(\d{3,6})', s)
    return f'IIF-{m.group(1)}' if m else s

def parse_date(v):
    s=iso(v)
    try: return datetime.strptime(s,'%Y-%m-%d').date() if s else None
    except: return None

def add_months(d, months):
    if not d: return None
    y=d.year+(d.month-1+months)//12; m=(d.month-1+months)%12+1
    return date(y,m,min(d.day,calendar.monthrange(y,m)[1]))

def month_key(d): return d.strftime('%Y-%m') if d else ''

def month_label(k):
    try: return datetime.strptime(k,'%Y-%m').strftime('%b %Y')
    except: return k

def get_settings():
    con=connect()
    d={r['key']:r['value'] for r in con.execute('SELECT key,value FROM settings')}; con.close()
    return {'issuer_limit':num(d.get('issuer_limit',2000000))}

def load_transactions():
    cached=getattr(g,'transactions_effective',None)
    if cached is not None: return cached
    base=read_parsed_source('transactions') if source_exists('transactions') else None
    if base is None:
        if not source_exists('transactions'): return []
        wb=openpyxl.load_workbook(read_source('transactions'),data_only=True,read_only=True)
        ws=wb[wb.sheetnames[0]]
        headers=[c.value for c in next(ws.iter_rows(min_row=2,max_row=2))]
        base=[]
        for vals in ws.iter_rows(min_row=3,values_only=True):
            if not any(v is not None for v in vals): continue
            r=dict(zip(headers,vals))
            base.append({
                'id': r.get('ID'), 'action': r.get('Action') or '', 'amount': num(r.get('Sum Involved(MYR)')),
                'previous_balance': num(r.get('Previous Balance(MYR)')), 'current_balance': num(r.get('Current Balance(MYR)')),
                'note': normalize_note_code(r.get('Note #')), 'transaction_no': r.get('Transaction No') or '',
                'status': r.get('Status') or '', 'date_raw': str(r.get('Date') or ''), 'date': iso(r.get('Date')), 'pay_via': r.get('Pay Via') or ''
            })
        wb.close()
    g.transactions_base=base
    con=connect()
    try: overrides={str(row['transaction_id']):row['loan_code'] for row in con.execute('SELECT transaction_id,loan_code FROM transaction_note_overrides')}
    finally: con.close()
    effective=[]
    for row in base:
        replacement=overrides.get(str(row.get('id')))
        effective.append({**row,'original_note':row.get('note',''),'note':replacement} if replacement else row)
    g.transactions_effective=effective
    return effective

def load_simulation():
    cached=getattr(g,'simulation_parsed',None)
    if cached is not None: return cached
    cached=read_parsed_source('simulation') if source_exists('simulation') else None
    if cached is not None:
        g.simulation_display=cached['display']
        g.simulation_parsed=(cached['portfolio'],cached['schedule'])
        return g.simulation_parsed
    if not source_exists('simulation'): return [], []
    wb=openpyxl.load_workbook(read_source('simulation'),data_only=True,read_only=True)
    ws=wb['Query result'] if 'Query result' in wb.sheetnames else wb[wb.sheetnames[0]]
    headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    display_headers=list(headers)
    while display_headers and display_headers[-1] is None: display_headers.pop()
    portfolio=[];display_rows=[]
    for vals in ws.iter_rows(min_row=2,values_only=True):
        if not vals or vals[0] is None or not vals[1]: continue
        display_rows.append(list(vals[:len(display_headers)]))
        r=dict(zip(headers,vals))
        expanded_profit='Total Gross Profit' in r
        gross_earned=num(r.get('Gross Profit Earned')) if expanded_profit else num(r.get('Gross Profit'))
        total_gross=num(r.get('Total Gross Profit')) if expanded_profit else num(r.get('Gross Profit'))
        item={
            'loan_code': normalize_note_code(r.get('Loan Code')),
            'issuer': r.get('Issuer Name') or '', 'note_name': r.get('Note Name') or '', 'product': r.get('Product Type') or '',
            'rating': r.get('CTOS/Payment Risk Rating **') or '', 'business': r.get('Business Description') or '',
            'investment_amount': num(r.get('Investment Amount')), 'note_size': num(r.get('Loan Note Size')),
            'exposure': num(r.get('Investment Exposure')), 'payment_type': r.get('Payment Type **') or '', 'term': r.get('Term') or '',
            'disbursal_date': iso(r.get('Disbursal Date')), 'final_date': iso(r.get('Final Repayment Date ')),
            'loan_status': r.get('Loan Status') or '', 'gross_pa': num(r.get('Interest Rate p.a (Gross)')),
            'gross_pm': num(r.get('Interest Rate p.m (Gross)')), 'net_pm': num(r.get('Interest Rate p.m (Net)')),
            'expected_repayment': num(r.get('Expected Repayment')), 'actual_repayment': num(r.get('Actual Repayment **')),
            'paid_principal': num(r.get('Paid Principal')), 'unpaid_principal': num(r.get('Unpaid Principal')),
            'gross_profit':total_gross,
            'net_profit': num(r.get('Net Profit')), 'paid_profit': num(r.get('Paid Profit')),
            'unpaid_profit': num(r.get('Unpaid Profit')),
            'late_profit': num(r.get('Late Payment Charges')) if expanded_profit else num(r.get('Late Profit')),
            'service_fee': num(r.get('Service Fee ')),
            'early_repayment': r.get('Early Repayment') or '', 'early_date': iso(r.get('Early Repayment Date')), 'remarks': r.get('Remarks') or ''
        }
        if expanded_profit:
            item.update({'gross_profit_earned':gross_earned,'sst':num(r.get('SST')),
                         'installment':r.get('Installment'),'report_format':'expanded-profit'})
        portfolio.append(item)
    schedule=[]
    if 'Sheet1' in wb.sheetnames:
        # A ReadOnlyWorksheet reparses its XML for every cell() call. Read the
        # small schedule table once so expanded reports stay within Worker CPU limits.
        table=list(wb['Sheet1'].iter_rows(min_row=4,values_only=True));months=[]
        heading=table[0] if table else ()
        for index in range(6,min(13,len(heading))):
            m=heading[index]
            if m: months.append((index,str(m)))
        for values in table[1:]:
            code=values[0] if values else None
            if not code: continue
            value=lambda index: values[index] if index<len(values) else None
            item={'loan_code':normalize_note_code(code),'monthly':num(value(1)),'gross_nett':value(2) or '',
                  'tenure':value(3) or '', 'final_date':iso(value(4)),'payment_type':value(5) or '', 'months':{}}
            for index,m in months:
                v=value(index)
                if isinstance(v,(int,float)): item['months'][m]=float(v)
                elif v: item['months'][m]=str(v)
            schedule.append(item)
    wb.close()
    g.simulation_display={'headers':display_headers,'rows':display_rows}
    g.simulation_parsed=(portfolio,schedule)
    return g.simulation_parsed


def load_projection_history():
    """Read realised historical receipt entries from the legacy MIDAS daily projection model.

    The legacy workbook models each day as: closing balance = opening balance + cash in - cash out.
    We use only the historical/realised portion to backfill old note repayment months; the live transaction
    ledger always wins when the same note/payment is present there.
    """
    cached=getattr(g,'projection_history_parsed',None)
    if cached is not None: return cached
    cached=read_parsed_source('projections') if source_exists('projections') else None
    if cached is not None:
        for event in cached.get('events',[]): event['date']=parse_date(event.get('date'))
        g.projection_history_parsed=cached
        return cached
    if not source_exists('projections'): return {'cutoff':None,'events':[]}
    wb=openpyxl.load_workbook(read_source('projections'),data_only=True,read_only=True)
    # Determine the last explicitly historical month from Sheet1's "as at" blocks.
    cutoff=None
    if 'Sheet1' in wb.sheetnames:
        ws1=wb['Sheet1']
        for row in ws1.iter_rows(values_only=True):
            for v in row:
                if isinstance(v,str):
                    m=re.search(r'as\s+at\s+(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?',v,re.I)
                    if m:
                        yy=int(m.group(3)) if m.group(3) else 2026
                        if yy<100: yy+=2000
                        try:
                            d=date(yy,int(m.group(2)),int(m.group(1)))
                            cutoff=max(cutoff,d) if cutoff else d
                        except: pass
    if cutoff is None: cutoff=date(2026,6,30)
    events=[]
    if 'Sheet2' in wb.sheetnames:
        ws=wb['Sheet2']
        # Date headings are on row 3; Cash In detail occupies rows 8:51 in this workbook.
        # Some Excel/OpenPyXL combinations can expose max_column/max_row as None
        # for a partially read or effectively empty worksheet.  Treat that source
        # as having no historical schedule instead of taking down the whole API.
        max_col = getattr(ws, 'max_column', 0) or 0
        max_row = getattr(ws, 'max_row', 0) or 0
        if max_col < 4 or max_row < 8:
            wb.close()
            result={'cutoff': cutoff.isoformat() if cutoff else None, 'events': []}
            g.projection_history_parsed=result
            return result
        # Parse the relevant range once. Random cell access on a read-only sheet
        # scans the underlying XML repeatedly and can exhaust Worker CPU time.
        table=list(ws.iter_rows(min_row=3,max_row=min(51,max_row),min_col=1,max_col=max_col,values_only=True))
        heading=table[0] if table else ()
        dates={index:parse_date(heading[index]) for index in range(3,len(heading))}
        current_label=''
        for values in table[5:]:
            label=values[1] if len(values)>1 else None
            typ=str(values[2] if len(values)>2 else '').strip().title()
            if typ not in ('Profit','Principal'): continue
            # Principal rows carry the note label; the following Profit row often carries only timing text.
            if typ=='Principal' and label: current_label=str(label)
            elif label and 'IIF' in str(label).upper(): current_label=str(label)
            code=normalize_note_code(current_label)
            if not code.startswith('IIF-'): continue
            for index,d in dates.items():
                if not d or d>cutoff: continue
                amt=values[index] if index<len(values) else None
                if isinstance(amt,(int,float)) and abs(float(amt))>0.005:
                    events.append({'loan_code':code,'type':typ,'date':d,'amount':float(amt),'source':'MIDAS Projections'})
    wb.close()
    result={'cutoff':cutoff.isoformat() if cutoff else None,'events':events}
    g.projection_history_parsed=result
    return result


def normalized_source(kind):
    """Parse one workbook into the compact representation used by calculations."""
    if kind=='transactions':
        rows=load_transactions()
        return getattr(g,'transactions_base',rows)
    if kind=='simulation':
        portfolio,schedule=load_simulation()
        return {'portfolio':portfolio,'schedule':schedule,'display':g.simulation_display}
    if kind=='projections':
        return load_projection_history()
    raise ValueError('Invalid source kind')

def ledger_payouts(transactions, code, action, exclude_special_profit=False):
    """Return successful payout totals by transaction date for a note."""
    vals={}
    for t in transactions:
        if t.get('note')!=code or t.get('action')!=action or str(t.get('status','')).upper()!='SUCCESSFUL':
            continue
        detail=str(t.get('transaction_no') or '').lower()
        if exclude_special_profit and ('deferred profit' in detail or 'tawidh' in detail or "ta'widh" in detail):
            continue
        d=parse_date(t.get('date'))
        if d: vals[d]=vals.get(d,0)+num(t.get('amount'))
    return sorted(vals.items())


def full_principal_settlement_date(r, transactions):
    """Detect the date on which cumulative principal payouts fully settled the MISB allocation."""
    target=max(0,num(r.get('investment_amount')))
    payouts=ledger_payouts(transactions,r['loan_code'],'Principal Payout')
    if not payouts: return None
    running=0
    tolerance=max(1.0,target*0.00001)
    for d,amt in payouts:
        running+=amt
        if target and running >= target-tolerance:
            return d
    # The simulation's completed/early-settlement flag is useful when rounding or legacy rows prevent exact equality.
    if str(r.get('loan_status') or '').lower()=='completed' and r.get('early_date'):
        return parse_date(r.get('early_date')) or payouts[-1][0]
    return None


def apply_early_settlement_to_profit_events(r, events, settlement):
    """Close future profit periods after an early full-principal settlement and insert the pro-rated final profit period.

    Pro-rata rule: outstanding principal × monthly rate × (days from last instalment date to settlement / 30).
    For Profit Only notes the outstanding principal is the original MISB allocation. For Bullet notes the period starts
    at disbursement. Equal-instalment notes use the scheduled outstanding principal after periods already due.
    """
    if not settlement: return events
    disb=parse_date(r.get('disbursal_date')); final=parse_date(r.get('final_date'))
    if not disb or not final or settlement >= final: return events
    ptype=str(r.get('payment_type') or '').lower(); rate=max(0,num(r.get('net_pm')))
    original=max(0,num(r.get('investment_amount')))
    if 'bullet' in ptype:
        days=max(0,(settlement-disb).days)
        expected=original*rate*days/30
        return [{'due_date':settlement,'expected':expected,'early_settlement':True,'original_due_date':final,'prorata_days':days}]

    kept=[dict(e) for e in events if e['due_date'] <= settlement]
    if any(e['due_date']==settlement for e in kept):
        kept[-1]['early_settlement']=True; kept[-1]['original_due_date']=kept[-1]['due_date']; kept[-1]['prorata_days']=30
        return kept
    previous_due=max([e['due_date'] for e in events if e['due_date'] < settlement], default=disb)
    next_due=min([e['due_date'] for e in events if e['due_date'] > settlement], default=final)
    days=max(0,(settlement-previous_due).days)
    outstanding=original
    if 'equal instalment' in ptype:
        try: term=max(1,int(float(r.get('term') or 1)))
        except: term=1
        periods_due=sum(1 for e in events if e['due_date'] <= previous_due)
        outstanding=max(0,original-(original/term)*periods_due)
    expected=outstanding*rate*days/30
    kept.append({'due_date':settlement,'expected':expected,'early_settlement':True,'original_due_date':next_due,'prorata_days':days})
    kept.sort(key=lambda e:e['due_date'])
    return kept


def match_profit_payouts(events, payouts):
    """Allocate ledger profit cash across scheduled periods, including multiple periods paid on one transaction date."""
    eps=0.005
    idx=0
    for pd,pool0 in payouts:
        pool=max(0,num(pool0))
        while pool>eps and idx<len(events):
            e=events[idx]
            expected=max(0,num(e.get('expected')))
            # Normal periods consume their scheduled amount. The final/early-settlement period consumes the remainder.
            if e.get('early_settlement') or idx==len(events)-1:
                actual=pool
            else:
                actual=min(pool, expected if expected>eps else pool)
            e.update({'paid':True,'paid_date':pd,'actual':actual})
            pool-=actual; idx+=1
    for i,e in enumerate(events):
        if 'paid' not in e: e.update({'paid':False,'paid_date':None,'actual':0})
    return events



def reconcile_profit_events(r, standard_events, settlement, payouts):
    """Reconcile scheduled profit with ledger cash. Early settlement reserves the settlement-date profit as
    the final pro-rated period, and determines the pro-rata start from the last instalment actually covered.
    """
    if not settlement:
        return match_profit_payouts([dict(e) for e in standard_events],payouts)
    disb=parse_date(r.get('disbursal_date')); final=parse_date(r.get('final_date'))
    if not disb or not final or settlement >= final:
        return match_profit_payouts([dict(e) for e in standard_events],payouts)

    before=[(d,a) for d,a in payouts if d < settlement]
    on_settlement=sum(a for d,a in payouts if d == settlement)
    after=[(d,a) for d,a in payouts if d > settlement]
    regular=[dict(e) for e in standard_events if e['due_date'] < settlement]
    regular=match_profit_payouts(regular,before)
    # Keep only regular periods actually satisfied before settlement. Any unsatisfied later schedule is extinguished.
    regular_paid=[e for e in regular if e.get('paid')]
    last_instalment=max([e['due_date'] for e in regular_paid],default=disb)
    ptype=str(r.get('payment_type') or '').lower()
    if 'bullet' in ptype: last_instalment=disb
    days=max(0,(settlement-last_instalment).days)
    outstanding=max(0,num(r.get('investment_amount')))
    if 'equal instalment' in ptype:
        try: term=max(1,int(float(r.get('term') or 1)))
        except: term=1
        outstanding=max(0,outstanding-(outstanding/term)*len(regular_paid))
    net_rate=max(0,num(r.get('net_pm'))); gross_rate=max(0,num(r.get('gross_pm')))
    early={'due_date':settlement,'expected':outstanding*net_rate*days/30,'expected_gross':outstanding*gross_rate*days/30,
           'early_settlement':True,'original_due_date':final,'prorata_days':days,'prorata_from':last_instalment}
    if on_settlement>0:
        early.update({'paid':True,'paid_date':settlement,'actual':on_settlement})
    elif after:
        # Allow a short processing delay after the principal settlement; the principal date remains the economic settlement date.
        pd,amt=after[0]
        if (pd-settlement).days <= 3: early.update({'paid':True,'paid_date':pd,'actual':amt})
        else: early.update({'paid':False,'paid_date':None,'actual':0})
    else:
        early.update({'paid':False,'paid_date':None,'actual':0})
    return regular_paid+[early]

def derived_profit_schedule(portfolio, transactions, projection_history=None):
    """Build profit periods from note terms, then reconcile them to the transaction ledger.
    A full early principal settlement closes all later periods and creates a final pro-rated profit period.
    """
    rows=[]; all_months=set()
    for r in portfolio:
        code=r['loan_code']; disb=parse_date(r.get('disbursal_date')); final=parse_date(r.get('final_date'))
        try: term=max(0,int(float(r.get('term') or 0)))
        except: term=0
        ptype=str(r.get('payment_type') or '').strip(); pl=ptype.lower(); events=[]
        if disb and term:
            if 'bullet' in pl:
                due=final or add_months(disb,term)
                expected=max(0,num(r.get('net_profit'))) or max(0,num(r.get('investment_amount'))*num(r.get('net_pm'))*term)
                if due: events.append({'due_date':due,'expected':expected})
            else:
                monthly=max(0,num(r.get('investment_amount'))*num(r.get('net_pm')))
                for i in range(1,term+1):
                    due=add_months(disb,i)
                    if final and i==term: due=final
                    if due: events.append({'due_date':due,'expected':monthly})
        elif final:
            events=[{'due_date':final,'expected':max(0,num(r.get('unpaid_profit')) or num(r.get('net_profit')))}]

        settlement=full_principal_settlement_date(r,transactions)
        payouts=ledger_payouts(transactions,code,'Profit Payout',True)
        # Backfill realised legacy periods only when the transaction ledger has no matching receipt.
        hist=[]
        if projection_history:
            hist=[(x['date'],x['amount']) for x in projection_history.get('events',[]) if x['loan_code']==code and x['type']=='Profit']
        if hist:
            ledger_dates={d for d,_ in payouts}
            for hd,ha in hist:
                if not any(abs((hd-ld).days)<=3 for ld in ledger_dates): payouts.append((hd,ha))
            payouts.sort(key=lambda z:z[0])
        events=reconcile_profit_events(r,events,settlement,payouts)
        for event in events:
            due=event['due_date']; all_months.add(month_key(due))
            if event['paid']:
                delta=(event['paid_date']-due).days
                event['status']='Paid'; event['timing']='early' if delta<0 else ('late' if delta>0 else 'on_time'); event['days_variance']=delta
            else:
                event['status']='Unpaid' if due < date.today() else 'Expected'; event['timing']='overdue' if due < date.today() else 'future'; event['days_variance']=None
        cells={}
        for e in events:
            k=month_key(e['due_date']); cells[k]={
                'due_date':e['due_date'].isoformat(),'expected':e['expected'],'paid':e['paid'],
                'paid_date':e['paid_date'].isoformat() if e['paid_date'] else '', 'actual':e['actual'],
                'status':e['status'],'timing':e['timing'],'days_variance':e['days_variance'],
                'early_settlement':bool(e.get('early_settlement')),
                'original_due_date':e.get('original_due_date').isoformat() if e.get('original_due_date') else '',
                'prorata_days':e.get('prorata_days'), 'prorata_from':e.get('prorata_from').isoformat() if e.get('prorata_from') else '',
                'expected_gross':e.get('expected_gross'),
                'source': ('MIDAS Projections' if projection_history and e.get('paid_date') and any(x['loan_code']==code and x['type']=='Profit' and x['date']==e.get('paid_date') for x in projection_history.get('events',[])) else ('Transaction ledger' if e.get('paid') else ''))
            }
        rows.append({'loan_code':code,'payment_type':ptype,'disbursal_date':r.get('disbursal_date'),'term':term,
                     'final_date':r.get('final_date'),'loan_status':r.get('loan_status'),'settlement_date':settlement.isoformat() if settlement else '', 'months':cells})
    months=sorted(all_months)
    return {'months':[{'key':k,'label':month_label(k)} for k in months], 'rows':rows}

def monitoring_map():
    con=connect()
    rows={r['loan_code']:dict(r) for r in con.execute('SELECT * FROM monitoring')}
    con.close(); return rows

def is_active(r): return bool(r['loan_status']) and r['loan_status'].lower() != 'completed'

def payment_marks_map():
    con=connect()
    rows=[dict(r) for r in con.execute('SELECT * FROM payment_marks ORDER BY expected_date,event_key')]
    con.close(); return {(r['loan_code'],r['event_key']):r for r in rows}

def cashflow_plan_rows():
    con=connect()
    rows=[dict(r) for r in con.execute('SELECT * FROM cashflow_plan ORDER BY flow_date,id')]
    con.close(); return rows

def note_payment_schedule(pf, schedule, transactions=None):
    marks=payment_marks_map(); transactions=transactions or []; out={}
    for r in pf:
        code=r['loan_code']; items=[]; disb=parse_date(r.get('disbursal_date')); final=parse_date(r.get('final_date'))
        try: term=max(0,int(float(r.get('term') or 0)))
        except: term=0
        ptype=str(r.get('payment_type') or '').lower(); profit_due=[]; principal_due=[]
        if disb and term:
            if 'bullet' in ptype:
                due=final or add_months(disb,term)
                if due:
                    profit_due=[{'due_date':due,'expected':max(0,num(r.get('net_profit'))) or max(0,num(r.get('investment_amount'))*num(r.get('net_pm'))*term)}]
                    principal_due=[{'due_date':due,'expected':max(0,num(r.get('investment_amount')))}]
            else:
                mp=max(0,num(r.get('investment_amount'))*num(r.get('net_pm')))
                for i in range(1,term+1):
                    due=final if (final and i==term) else add_months(disb,i)
                    if due: profit_due.append({'due_date':due,'expected':mp})
                if 'equal instalment' in ptype:
                    princ=max(0,num(r.get('investment_amount'))/term) if term else 0
                    principal_due=[{'due_date':x['due_date'],'expected':princ} for x in profit_due]
                elif final:
                    principal_due=[{'due_date':final,'expected':max(0,num(r.get('investment_amount')))}]
        elif final:
            profit_due=[{'due_date':final,'expected':max(0,num(r.get('unpaid_profit')) or num(r.get('net_profit')))}]
            principal_due=[{'due_date':final,'expected':max(0,num(r.get('unpaid_principal')) or num(r.get('investment_amount')))}]

        settlement=full_principal_settlement_date(r,transactions)
        profit_due=reconcile_profit_events(r,profit_due,settlement,ledger_payouts(transactions,code,'Profit Payout',True))

        # Once principal is fully settled, no later principal instalments remain receivable.
        if settlement and final and settlement < final:
            if 'equal instalment' in ptype:
                principal_due=[x for x in principal_due if x['due_date'] < settlement]
                paid_before=sum(a for d,a in ledger_payouts(transactions,code,'Principal Payout') if d < settlement)
                remaining=max(0,num(r.get('investment_amount'))-paid_before)
                principal_due.append({'due_date':settlement,'expected':remaining,'early_settlement':True,'original_due_date':final})
            else:
                principal_due=[{'due_date':settlement,'expected':max(0,num(r.get('investment_amount'))),'early_settlement':True,'original_due_date':final}]

        # Allocate principal cash sequentially; supports split transactions on the same day.
        pp=ledger_payouts(transactions,code,'Principal Payout'); pi=0
        for pd,pool0 in pp:
            pool=max(0,num(pool0))
            while pool>0.005 and pi<len(principal_due):
                e=principal_due[pi]; exp=max(0,num(e.get('expected')))
                actual=pool if e.get('early_settlement') or pi==len(principal_due)-1 else min(pool,exp if exp>0.005 else pool)
                e.update({'paid':True,'paid_date':pd,'actual':actual}); pool-=actual; pi+=1
        for e in principal_due:
            if 'paid' not in e: e.update({'paid':False,'paid_date':None,'actual':0})

        for typ,due_list in [('Profit',profit_due),('Principal',principal_due)]:
            for idx,e in enumerate(due_list):
                due=e['due_date']; expected=e['expected']; early=bool(e.get('early_settlement'))
                final_principal=(typ=='Principal' and (len(due_list)==1 or early))
                key=('principal-final' if final_principal else f'{typ.lower()}-{due.strftime("%Y-%m")}')
                if early: key=f'{typ.lower()}-early-{due.isoformat()}'
                if early and typ=='Profit':
                    label=f'Early settlement profit ({e.get("prorata_days",0)} days pro-rated)'
                elif early and typ=='Principal': label='Early principal settlement'
                else: label=('Final principal repayment' if final_principal else f'{due.strftime("%b %Y")} scheduled {typ.lower()}')
                base={'loan_code':code,'event_key':key,'event_type':typ,'label':label,'expected_date':due.isoformat(),
                      'expected_amount':expected,'is_paid':1 if e.get('paid') else 0,'paid_date':e['paid_date'].isoformat() if e.get('paid_date') else '',
                      'paid_amount':e.get('actual',0),'notes':'Auto-matched from transaction ledger' if e.get('paid') else '',
                      'early_settlement':early,'original_due_date':e.get('original_due_date').isoformat() if e.get('original_due_date') else '',
                      'prorata_days':e.get('prorata_days'),'prorata_from':e.get('prorata_from').isoformat() if e.get('prorata_from') else '',
                      'expected_gross':e.get('expected_gross')}
                manual=marks.get((code,key))
                if not manual and typ=='Profit' and not early:
                    manual=marks.get((code,'profit-'+due.strftime('%B').lower()))
                if manual and not e.get('paid'): base.update(manual)
                elif manual and e.get('paid') and manual.get('notes'): base['notes']=manual.get('notes')
                items.append(base)
        known={x['event_key'] for x in items}
        for (mc,mk),m in marks.items():
            if mc==code and mk not in known and not (mk.startswith('profit-') and any(x['event_type']=='Profit' and x.get('paid_date')==m.get('paid_date') for x in items)):
                items.append(m)
        items.sort(key=lambda x:(x.get('expected_date') or '9999-99-99',x.get('event_type') or ''))
        out[code]=items
    return out

def build_cash_projection(current_cash, last_tx_date, payment_schedule):
    events=[]
    for code,items in payment_schedule.items():
        for x in items:
            if x.get('event_type')=='Principal' and x.get('expected_date') and not x.get('is_paid'):
                events.append({'date':x['expected_date'],'type':'Receivable','category':'Principal','loan_code':code,'issuer':'','amount':num(x.get('expected_amount')),'notes':x.get('label','')})
    for x in cashflow_plan_rows():
        sign=-1 if x['flow_type']=='Allocation' else 1
        events.append({'date':x['flow_date'],'type':x['flow_type'],'category':x['flow_type'],'loan_code':x['loan_code'],'issuer':x['issuer'],'amount':sign*num(x['amount']),'notes':x['notes'],'id':x['id']})
    events=[x for x in events if x.get('date') and (not last_tx_date or x['date']>last_tx_date)]
    events.sort(key=lambda x:(x['date'],0 if x['amount']<0 else 1))
    bal=current_cash; timeline=[]
    for x in events:
        bal+=x['amount']; y=dict(x); y['projected_balance']=bal; timeline.append(y)
    return timeline


def build_reconciliation(transactions, portfolio, tolerance=0.01):
    """Compare note-level ledger movements with the imported simulation snapshot."""
    ledger={}
    for row in transactions:
        if str(row.get('status','')).upper()!='SUCCESSFUL': continue
        code=row.get('note') or ''
        if not code.startswith('IIF-'): continue
        item=ledger.setdefault(code,{'committed':0.0,'principal_paid':0.0,'profit_paid':0.0})
        action=row.get('action'); amount=max(0,num(row.get('amount')))
        if action=='Investment Committed': item['committed']+=amount
        elif action=='Principal Payout': item['principal_paid']+=amount
        elif action=='Profit Payout': item['profit_paid']+=amount
    simulation={row['loan_code']:row for row in portfolio if row.get('loan_code')}
    rows=[]
    for code in sorted(set(ledger)|set(simulation)):
        led=ledger.get(code,{'committed':0.0,'principal_paid':0.0,'profit_paid':0.0})
        sim=simulation.get(code)
        ledger_outstanding=max(0,led['committed']-led['principal_paid'])
        issues=[]
        if not sim: issues.append('Missing from simulation')
        elif code not in ledger: issues.append('No ledger allocation')
        investment_diff=led['committed']-num(sim.get('investment_amount')) if sim else led['committed']
        principal_diff=led['principal_paid']-num(sim.get('paid_principal')) if sim else led['principal_paid']
        profit_diff=led['profit_paid']-num(sim.get('paid_profit')) if sim else led['profit_paid']
        outstanding_diff=ledger_outstanding-num(sim.get('unpaid_principal')) if sim else ledger_outstanding
        for label,value in (('Investment',investment_diff),('Principal paid',principal_diff),
                            ('Profit paid',profit_diff),('Outstanding principal',outstanding_diff)):
            if abs(value)>tolerance: issues.append(f'{label} differs')
        rows.append({'loan_code':code,'issuer':sim.get('issuer','') if sim else '',
                     'ledger_committed':led['committed'],'simulation_investment':num(sim.get('investment_amount')) if sim else 0,
                     'ledger_principal_paid':led['principal_paid'],'simulation_principal_paid':num(sim.get('paid_principal')) if sim else 0,
                     'ledger_profit_paid':led['profit_paid'],'simulation_profit_paid':num(sim.get('paid_profit')) if sim else 0,
                     'ledger_outstanding':ledger_outstanding,'simulation_outstanding':num(sim.get('unpaid_principal')) if sim else 0,
                     'investment_difference':investment_diff,'principal_difference':principal_diff,
                     'profit_difference':profit_diff,'outstanding_difference':outstanding_diff,
                     'issues':issues,'status':'Matched' if not issues else 'Review'})
    mismatches=[row for row in rows if row['status']=='Review']
    return {'rows':rows,'matched':len(rows)-len(mismatches),'review':len(mismatches),
            'missing_from_simulation':sum('Missing from simulation' in row['issues'] for row in rows),
            'outstanding_difference':sum(row['outstanding_difference'] for row in rows)}


def build_transaction_matching(transactions, portfolio):
    known={row['loan_code'] for row in portfolio if row.get('loan_code')}
    relevant={'Investment Committed','Principal Payout','Profit Payout'}
    rows=[]
    for row in transactions:
        if str(row.get('status','')).upper()!='SUCCESSFUL' or row.get('action') not in relevant: continue
        note=row.get('note') or ''
        if note not in known or row.get('original_note') is not None:
            rows.append({'id':row.get('id'),'date':row.get('date'),'action':row.get('action'),
                         'amount':row.get('amount'),'note':note,
                         'original_note':row.get('original_note',note),
                         'overridden':row.get('original_note') is not None,
                         'reason':'Manual assignment' if row.get('original_note') is not None else ('Missing note number' if not note.startswith('IIF-') else 'Note not in simulation')})
    return {'rows':rows,'unresolved':sum(not row['overridden'] for row in rows),
            'overridden':sum(row['overridden'] for row in rows),
            'candidates':[{'loan_code':row['loan_code'],'issuer':row.get('issuer',''),'note_name':row.get('note_name','')} for row in portfolio]}

def build_payload():
    tx=load_transactions(); pf,schedule=load_simulation(); projection_history=load_projection_history();
    if not isinstance(projection_history, dict):
        projection_history={'cutoff':None,'events':projection_history if isinstance(projection_history,list) else []}
    mon=monitoring_map(); settings=get_settings(); issuer_limit=settings['issuer_limit']; pay_schedule=note_payment_schedule(pf,schedule,tx)
    profit_schedule=derived_profit_schedule(pf,tx,projection_history)
    for r in pf:
        r['monitoring']=mon.get(r['loan_code'], {'status':'Normal','owner':'','next_action':'','review_date':'','notes':''})
        r['payment_schedule']=pay_schedule.get(r['loan_code'],[])
    successful=[r for r in tx if str(r['status']).upper()=='SUCCESSFUL']
    latest_cash=successful[0]['current_balance'] if successful else (tx[0]['current_balance'] if tx else 0)
    invested=sum(r['amount'] for r in successful if r['action']=='Investment Committed')
    principal_paid=sum(r['amount'] for r in successful if r['action']=='Principal Payout')
    profit_paid=sum(r['amount'] for r in successful if r['action']=='Profit Payout')
    outstanding_principal=max(0, invested-principal_paid)
    deployed=outstanding_principal
    today_d=date.today(); q_start_month=((today_d.month-1)//3)*3+1; q_start=date(today_d.year,q_start_month,1)
    q_profit=sum(r['amount'] for r in successful if r['action']=='Profit Payout' and parse_date(r.get('date')) and q_start<=parse_date(r.get('date'))<=today_d)
    return_this_quarter=q_profit/10000000
    total_return=profit_paid/10000000
    active=[r for r in pf if is_active(r)]
    sim_outstanding=sum(max(0,r['unpaid_principal']) for r in active)
    unpaid_profit=sum(max(0,r['unpaid_profit']) for r in active)
    projected_net_profit=sum(max(0,r['net_profit']) for r in active)
    paid_profit_sim=sum(max(0,r['paid_profit']) for r in pf)
    service_fee=sum(max(0,r['service_fee']) for r in pf)
    total_value=latest_cash+deployed
    weighted_monthly=(sum(max(0,r['unpaid_principal'])*max(0,r['net_pm']) for r in active)/sim_outstanding) if sim_outstanding else 0
    today=date.today().isoformat()
    due=[]; daily={}
    for r in active:
        if r['final_date'] and re.match(r'^\d{4}-\d{2}-\d{2}$',r['final_date']):
            d=datetime.strptime(r['final_date'],'%Y-%m-%d').date(); days=(d-date.today()).days
            rr=dict(r); rr['days_to_due']=days; rr['expected_principal']=max(0,r['unpaid_principal']); due.append(rr)
            x=daily.setdefault(r['final_date'],{'date':r['final_date'],'principal':0,'notes':0,'items':[]})
            x['principal']+=max(0,r['unpaid_principal']); x['notes']+=1; x['items'].append({'loan_code':r['loan_code'],'issuer':r['issuer'],'amount':max(0,r['unpaid_principal']),'payment_type':r['payment_type']})
    due.sort(key=lambda x:x['days_to_due'])
    receivables=sorted(daily.values(),key=lambda x:x['date'])
    issuer={}
    for r in active:
        k=r['issuer'] or 'Unknown Issuer'; d=issuer.setdefault(k,{'issuer':k,'outstanding':0,'unpaid_profit':0,'notes':0,'rating':r['rating']})
        d['outstanding']+=max(0,r['unpaid_principal']); d['unpaid_profit']+=max(0,r['unpaid_profit']); d['notes']+=1
    issuers=[]
    for d in issuer.values():
        d['limit']=issuer_limit; d['headroom']=max(0,issuer_limit-d['outstanding']); d['over_limit']=max(0,d['outstanding']-issuer_limit); d['utilization']=d['outstanding']/issuer_limit if issuer_limit else 0
        d['allocatable_now']=min(latest_cash,d['headroom']); issuers.append(d)
    issuers.sort(key=lambda x:x['outstanding'],reverse=True)
    action_totals={}
    for r in successful: action_totals[r['action']]=action_totals.get(r['action'],0)+r['amount']
    balance_breakdown=[
        {'label':'Cash available','amount':latest_cash,'note':'Latest successful wallet balance'},
        {'label':'Total investment committed','amount':invested,'note':'Cumulative successful allocations since inception'},
        {'label':'Outstanding principal','amount':outstanding_principal,'note':'Investment committed less principal payouts'},
        {'label':'Unpaid net profit','amount':unpaid_profit,'note':'Expected future profit from active simulation notes'},
    ]
    last_tx_date=successful[0]['date'] if successful else (tx[0]['date'] if tx else None)
    try: sim_as_of=source_rows().get('simulation',{}).get('as_of','Unknown')
    except: sim_as_of='Unknown'
    projection=build_cash_projection(latest_cash,last_tx_date,pay_schedule)
    reconciliation=build_reconciliation(tx,pf)
    transaction_matching=build_transaction_matching(tx,pf)
    from allocations import allocation_map
    all_allocation_records=allocation_map()
    approved_allocations={code for code,record in all_allocation_records.items()
                          if record.get('approval_status') in ('Approved','Disbursed')}
    allocation_workflow={status:sum(record.get('approval_status','Draft')==status for record in all_allocation_records.values())
                         for status in ('Draft','Ready for Approval','Approved','Disbursed','Cancelled')}
    pending_allocations=[
        {'loan_code':row['loan_code'],'allocated':row['ledger_committed']}
        for row in reconciliation['rows']
        if 'Missing from simulation' in row['issues']
        and row['ledger_committed']>0
        and row['loan_code'] not in approved_allocations
    ]
    return {'settings':settings,'summary':{'cash':latest_cash,'ledger_deployed':deployed,'total_investment_committed':invested,'outstanding_principal':outstanding_principal,'ledger_profit_paid':profit_paid,'quarter_profit_collected':q_profit,'return_this_quarter':return_this_quarter,'total_return':total_return,'implied_fund_value':total_value,
                       'simulation_outstanding':sim_outstanding,'simulation_unpaid_profit':unpaid_profit,'active_notes':len(active),
                       'transaction_count':len(tx),'portfolio_count':len(pf),'last_transaction_date':last_tx_date,'simulation_as_of':sim_as_of,
                       'portfolio_gap':deployed-sim_outstanding,'projected_net_profit':projected_net_profit,'paid_profit_simulation':paid_profit_sim,
                       'weighted_monthly_net_rate':weighted_monthly,'annualized_net_rate':weighted_monthly*12,'service_fee':service_fee,
                       'investable_cash':latest_cash},
            'transactions':tx,'portfolio':pf,'schedule':schedule,'profit_schedule':profit_schedule,'due':due,'receivables':receivables,'issuers':issuers,
            'balance_breakdown':balance_breakdown,'action_totals':action_totals,'today':today,'cash_projection':projection,'cashflow_plan':cashflow_plan_rows(),
            'reconciliation':reconciliation,'transaction_matching':transaction_matching,'pending_allocations':pending_allocations,
            'allocation_workflow':allocation_workflow,
            'projection_history':{'cutoff':projection_history.get('cutoff'),'event_count':len(projection_history.get('events',[]))}}



@app.route('/api/data')
def data(): return jsonify(build_payload())

@app.route('/api/settings', methods=['POST'])
def save_settings():
    d=request.get_json(force=True) or {}; limit=num(d.get('issuer_limit'))
    if limit <= 0: return jsonify({'error':'Issuer limit must be greater than zero'}),400
    con=connect(); con.execute("INSERT INTO settings(key,value) VALUES('issuer_limit',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(str(limit),)); con.commit(); con.close()
    audit_event('Updated issuer limit','settings','issuer_limit',{'value':limit})
    return jsonify({'ok':True})


@app.route('/api/transaction-note-override/<transaction_id>',methods=['POST','DELETE'])
def transaction_note_override(transaction_id):
    con=connect()
    try:
        if request.method=='DELETE':
            con.execute('DELETE FROM transaction_note_overrides WHERE transaction_id=?',(transaction_id,));con.commit()
            audit_event('Removed assignment','transaction',transaction_id)
            return jsonify({'ok':True})
        code=normalize_note_code((request.get_json() or {}).get('loan_code'))
        portfolio,_=load_simulation()
        if code not in {row['loan_code'] for row in portfolio}:
            return jsonify({'error':'Choose a note from the current simulation.'}),400
        con.execute('''INSERT INTO transaction_note_overrides(transaction_id,loan_code) VALUES(?,?)
          ON CONFLICT(transaction_id) DO UPDATE SET loan_code=excluded.loan_code,updated_at=CURRENT_TIMESTAMP''',(transaction_id,code))
        con.commit()
    finally:
        con.close()
    audit_event('Assigned note','transaction',transaction_id,{'loan_code':code})
    return jsonify({'ok':True,'transaction_id':transaction_id,'loan_code':code})


@app.get('/api/audit-events')
def audit_events():
    return jsonify(audit_rows(request.args.get('limit',200)))

@app.route('/api/monitoring/<loan_code>', methods=['POST'])
def save_monitoring(loan_code):
    d=request.get_json(force=True) or {}
    vals=(loan_code,d.get('status','Normal'),d.get('owner',''),d.get('next_action',''),d.get('review_date',''),d.get('notes',''))
    con=connect()
    con.execute('''INSERT INTO monitoring(loan_code,status,owner,next_action,review_date,notes,updated_at)
      VALUES(?,?,?,?,?,?,CURRENT_TIMESTAMP) ON CONFLICT(loan_code) DO UPDATE SET status=excluded.status, owner=excluded.owner,
      next_action=excluded.next_action, review_date=excluded.review_date, notes=excluded.notes, updated_at=CURRENT_TIMESTAMP''',vals)
    con.commit(); con.close(); audit_event('Updated monitoring','note',loan_code,{'status':d.get('status','Normal')}); return jsonify({'ok':True})

@app.route('/api/payment-mark', methods=['POST'])
def save_payment_mark():
    d=request.get_json(force=True) or {}
    code=str(d.get('loan_code') or '').strip(); key=str(d.get('event_key') or '').strip()
    if not code or not key: return jsonify({'error':'loan_code and event_key are required'}),400
    vals=(code,key,d.get('event_type','Other'),d.get('label',''),d.get('expected_date',''),num(d.get('expected_amount')),1 if d.get('is_paid') else 0,d.get('paid_date',''),num(d.get('paid_amount')),d.get('notes',''))
    con=connect()
    con.execute("""INSERT INTO payment_marks(loan_code,event_key,event_type,label,expected_date,expected_amount,is_paid,paid_date,paid_amount,notes,updated_at)
      VALUES(?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP) ON CONFLICT(loan_code,event_key) DO UPDATE SET event_type=excluded.event_type,label=excluded.label,expected_date=excluded.expected_date,expected_amount=excluded.expected_amount,is_paid=excluded.is_paid,paid_date=excluded.paid_date,paid_amount=excluded.paid_amount,notes=excluded.notes,updated_at=CURRENT_TIMESTAMP""",vals)
    con.commit(); con.close(); return jsonify({'ok':True})

@app.route('/api/cashflow-plan', methods=['POST'])
def add_cashflow_plan():
    d=request.get_json(force=True) or {}; typ=d.get('flow_type')
    if typ not in ('Allocation','Receivable'): return jsonify({'error':'Flow type must be Allocation or Receivable'}),400
    if not d.get('flow_date') or num(d.get('amount'))<=0: return jsonify({'error':'Date and positive amount are required'}),400
    con=connect(); cur=con.execute('INSERT INTO cashflow_plan(flow_date,flow_type,loan_code,issuer,amount,notes) VALUES(?,?,?,?,?,?)',(d['flow_date'],typ,d.get('loan_code',''),d.get('issuer',''),num(d['amount']),d.get('notes',''))); con.commit(); i=cur.lastrowid; con.close(); audit_event('Added planned cash flow','cashflow',i,{'type':typ,'date':d['flow_date'],'amount':num(d['amount'])}); return jsonify({'ok':True,'id':i})

@app.route('/api/cashflow-plan/<int:item_id>', methods=['DELETE'])
def delete_cashflow_plan(item_id):
    con=connect(); con.execute('DELETE FROM cashflow_plan WHERE id=?',(item_id,)); con.commit(); con.close(); audit_event('Deleted planned cash flow','cashflow',item_id); return jsonify({'ok':True})

@app.route('/api/upload/<kind>', methods=['POST'])
def upload(kind):
    from uploads import inspect_workbook
    if kind not in ('transactions','simulation','projections'):
        return jsonify({'error':'Invalid file type'}),400
    f=request.files.get('file')
    if not f or not f.filename.lower().endswith('.xlsx'):
        return jsonify({'error':'Please upload an .xlsx file'}),400
    data=f.read()
    as_of=request.form.get('as_of') or date.today().isoformat()
    if not parse_date(as_of):
        return jsonify({'error':'As-of date must be YYYY-MM-DD'}),400
    parsed=None
    try:
        current=None
        if kind=='transactions':
            loaded=load_transactions();current=list(getattr(g,'transactions_base',loaded))
        elif kind=='simulation':
            current=list(load_simulation()[0])
        quality=inspect_workbook(kind,data,current)
        if not quality.get('can_import'):
            return jsonify({'error':'Workbook failed data-quality checks. Validate it to review the issues.',
                            'issues':quality.get('issues',[]),'stats':quality.get('stats',{})}),422
        comparison=quality.get('comparison') or {};allow_regression=request.form.get('allow_regression') in ('1','true','yes')
        if comparison.get('requires_acknowledgement') and not allow_regression:
            return jsonify({'error':'This workbook would remove or roll back live source data. Review and acknowledge the regression before importing.',
                            'comparison':comparison}),409
        # Parse only the candidate source before publishing it. The normalized
        # result is stored in D1 so ordinary dashboard requests never reopen XLSX.
        for key in ('transactions_effective','transactions_base','simulation_parsed','simulation_display'):
            g.pop(key,None)
        source_rows()
        g.sources[kind]={'kind':kind,'object_key':'candidate','as_of':as_of}
        if 'source_bytes' not in g: g.source_bytes={}
        g.source_bytes[kind]=data
        parsed=normalized_source(kind)
    except Exception:
        return jsonify({'error':'Workbook could not be read. Check its template, dates and amounts.'}),400
    finally:
        g.pop('sources',None)
        g.pop('source_bytes',None)
    backup=save_source(kind,data,as_of,f.filename)
    current=source_rows()[kind]
    save_parsed_source(kind,current['object_key'],parsed)
    audit_event('Imported workbook','source',kind,{'as_of':as_of,'object_key':current['object_key'],
        'backup':backup or '','regression_acknowledged':bool((quality.get('comparison') or {}).get('requires_acknowledgement'))})
    return jsonify({'ok':True,'backup':backup or 'First import; no previous version','validation':quality})


@app.post('/api/upload/validate/<kind>')
def validate_upload(kind):
    from uploads import inspect_workbook
    if kind not in ('transactions','simulation','projections'):
        return jsonify({'error':'Invalid file type'}),400
    f=request.files.get('file')
    if not f or not f.filename.lower().endswith('.xlsx'):
        return jsonify({'error':'Please upload an .xlsx file'}),400
    try:
        current=None
        if kind=='transactions':
            loaded=load_transactions();current=getattr(g,'transactions_base',loaded)
        elif kind=='simulation':
            current=load_simulation()[0]
        return jsonify(inspect_workbook(kind,f.read(),current))
    except Exception as error:
        return jsonify({'error':str(error) or 'Workbook could not be validated.'}),400


@app.get('/api/source-versions/<kind>')
def list_source_versions(kind):
    if kind not in ('transactions','simulation','projections'):
        return jsonify({'error':'Invalid source kind'}),400
    return jsonify(source_versions(kind))


@app.get('/api/source-versions/<kind>/download')
def download_source_version(kind):
    if kind not in ('transactions','simulation','projections'):
        return jsonify({'error':'Invalid source kind'}),400
    from storage import read_source_version
    from werkzeug.utils import secure_filename
    object_key=str(request.args.get('object_key') or '')
    try:
        stream,version=read_source_version(kind,object_key)
        fallback=f"{kind}_{version.get('as_of') or 'undated'}_{object_key.split('/')[-1][:8]}.xlsx"
        filename=secure_filename(version.get('filename') or fallback) or fallback
        if not filename.lower().endswith('.xlsx'): filename+='.xlsx'
        return send_file(stream,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True,download_name=filename)
    except ValueError as error:
        return jsonify({'error':str(error)}),404


@app.post('/api/source-versions/<kind>/restore')
def restore_source_version(kind):
    if kind not in ('transactions','simulation','projections'):
        return jsonify({'error':'Invalid source kind'}),400
    body=request.get_json() or {};object_key=str(body.get('object_key') or '')
    previous=source_rows().get(kind,{}).copy()
    try:
        as_of=select_source_version(kind,object_key,str(body.get('as_of') or ''))
        for key in ('transactions_parsed','simulation_parsed','simulation_display','projection_history_parsed'):
            g.pop(key,None)
        parsed=normalized_source(kind)
        save_parsed_source(kind,object_key,parsed)
        audit_event('Restored workbook','source',kind,{'object_key':object_key,'as_of':as_of})
        return jsonify({'ok':True,'kind':kind,'as_of':as_of})
    except Exception as error:
        if previous.get('object_key'):
            select_source_version(kind,previous['object_key'],previous.get('as_of',''))
        return jsonify({'error':str(error) or 'The source version could not be restored.'}),400


@app.post('/api/admin/rebuild-source-cache/<kind>')
def rebuild_source_cache(kind):
    """One-time/backfill route for workbooks imported before parsed D1 storage."""
    if kind not in ('transactions','simulation','projections'):
        return jsonify({'error':'Invalid source kind'}),400
    row=source_rows().get(kind)
    if not row: return jsonify({'error':'Source is not imported'}),404
    parsed=normalized_source(kind)
    save_parsed_source(kind,row['object_key'],parsed)
    return jsonify({'ok':True,'kind':kind})


SIMULATION_COMPARE_FIELDS=('Loan Status','Investment Amount','Paid Principal','Unpaid Principal',
                           'Paid Profit','Unpaid Profit','Late Profit','SST')


def compact_simulation_snapshot(headers, rows):
    positions={name:headers.index(name) for name in ('Loan Code',)+SIMULATION_COMPARE_FIELDS if name in headers}
    if 'Loan Code' not in positions: return {}
    result={}
    for row in rows:
        if len(row)<=positions['Loan Code']: continue
        code=normalize_note_code(row[positions['Loan Code']])
        if not code: continue
        result[code]={name:row[index] if index<len(row) else None for name,index in positions.items() if name!='Loan Code'}
    return result


def simulation_comparison(headers, rows, prior_id=None):
    from storage import simulation_report_run,simulation_report_runs
    if prior_id is not None:
        prior=simulation_report_run(prior_id)
    else:
        history=simulation_report_runs(1);prior=history[0] if history else None
    if not prior:
        return {'has_prior':False,'requested_id':prior_id,'new_notes':[],'removed_notes':[],'changed_notes':[]}
    before=prior.get('snapshot') or {};current=compact_simulation_snapshot(headers,rows)
    new_notes=sorted(set(current)-set(before));removed_notes=sorted(set(before)-set(current));changed=[]
    for code in sorted(set(current)&set(before)):
        fields=[]
        for field in SIMULATION_COMPARE_FIELDS:
            old=before[code].get(field);new=current[code].get(field)
            if isinstance(old,(int,float)) or isinstance(new,(int,float)):
                if abs(num(old)-num(new))>.005: fields.append({'field':field,'before':old,'after':new})
            elif old!=new: fields.append({'field':field,'before':old,'after':new})
        if fields: changed.append({'loan_code':code,'fields':fields})
    return {'has_prior':True,'prior_id':prior['id'],'prior_as_of':prior['as_of'],
            'prior_created_at':prior['created_at'],'new_notes':new_notes,
            'removed_notes':removed_notes,'changed_notes':changed}


def simulation_update_snapshot(as_of=None, include_drafts=False, compare_to=None):
    """Build copy/export values from the current simulation template, updated from successful ledger transactions.

    Only transaction-driven servicing fields are changed. Descriptive fields, remarks, note names and other template content
    remain exactly as they are in the imported simulation workbook.
    """
    cutoff=parse_date(as_of) if as_of else None
    tx=[t for t in load_transactions() if str(t.get('status','')).upper()=='SUCCESSFUL' and (not cutoff or (parse_date(t.get('date')) and parse_date(t.get('date'))<=cutoff))]
    portfolio,_=load_simulation(); by_code={r['loan_code']:r for r in portfolio}
    updates={}
    for code,r in by_code.items():
        note_tx=[t for t in tx if t.get('note')==code]
        inv=max(0,num(r.get('investment_amount')))
        paid_principal=sum(num(t.get('amount')) for t in note_tx if t.get('action')=='Principal Payout')
        tawidh=sum(num(t.get('amount')) for t in note_tx if t.get('action')=='Profit Payout' and ('tawidh' in str(t.get('transaction_no') or '').lower() or "ta'widh" in str(t.get('transaction_no') or '').lower()))
        profit_tx=[t for t in note_tx if t.get('action')=='Profit Payout']
        all_profit=sum(num(t.get('amount')) for t in profit_tx)
        paid_profit=all_profit if r.get('report_format')=='expanded-profit' else all_profit-tawidh
        unpaid_principal=max(0,inv-paid_principal)
        settlement=None; running=0.0
        principal_events=[]
        for t in note_tx:
            if t.get('action')!='Principal Payout': continue
            d=parse_date(t.get('date'))
            if d: principal_events.append((d,num(t.get('amount'))))
        principal_events.sort(key=lambda z:z[0])
        tol=max(1.0,inv*0.00001)
        for d,a in principal_events:
            running+=a
            if inv and running>=inv-tol:
                settlement=d; break
        final=parse_date(r.get('final_date'))
        completed=bool(inv and paid_principal>=inv-tol)
        early=bool(completed and settlement and final and settlement<final)
        original_net=max(0,num(r.get('net_profit')))
        original_gross=max(0,num(r.get('gross_profit')))
        expanded=r.get('report_format')=='expanded-profit'
        # SST is transaction-driven: it applies only to profit payouts dated
        # 1 August 2026 onward, with the same per-transaction rounding as statements.
        from statements import profit_components
        sst=sum(float(profit_components(t.get('amount'),parse_date(t.get('date')))[2])
                for t in profit_tx if parse_date(t.get('date')))
        gross_earned=max(0,num(r.get('gross_profit_earned')))
        late_charge=max(0,num(r.get('late_profit')))
        if completed and expanded:
            expected_net=paid_profit;unpaid_profit=0.0
            gross_profit=(expected_net+sst)/0.8 if expected_net or sst else 0
            if early:
                gross_earned=gross_profit;late_charge=0
            else:
                late_charge=max(0,gross_profit-gross_earned)
            service_fee=max(0,gross_profit*.2)
        elif completed:
            expected_net=paid_profit;unpaid_profit=0.0
            gross_profit=(expected_net/0.8 if expected_net else original_gross)
            service_fee=max(0,gross_profit-expected_net)
        else:
            gross_profit=original_gross;service_fee=max(0,num(r.get('service_fee')))
            expected_net=max(0,gross_profit-service_fee-sst) if expanded else original_net
            unpaid_profit=max(0,expected_net-paid_profit)
        updates[code]={
            '_expanded':expanded,
            'Loan Status':'Completed' if completed else r.get('loan_status',''),
            'Actual Repayment **':paid_principal+paid_profit,
            'Paid Principal':paid_principal,
            'Unpaid Principal':unpaid_principal,
            'Gross Profit':gross_profit,
            'Gross Profit Earned':gross_earned,
            'Total Gross Profit':gross_profit,
            'Net Profit':expected_net,
            'Paid Profit':paid_profit,
            'Unpaid Profit':unpaid_profit,
            'Late Profit':tawidh,
            'Late Payment Charges':late_charge,
            'Service Fee ':service_fee,
            'SST':sst,
            'Installment':r.get('installment',''),
            'Early Repayment':'Early Repayment of Principal' if early else (r.get('early_repayment','') if not completed else ''),
            'Early Repayment Date':settlement.isoformat() if early and settlement else (r.get('early_date') if not completed else None),
        }
    # load_simulation caches the displayed rows for this request, avoiding a
    # second OpenPyXL parse of the same workbook on CPU-limited Workers.
    displayed=g.simulation_display;headers=displayed['headers']
    rows=[]; template_codes=set();ledger_updated_rows=[]
    for vals in displayed['rows']:
        arr=list(vals); code=normalize_note_code(arr[1]); template_codes.add(code)
        up=updates.get(code,{})
        for i,h in enumerate(headers):
            if h in up:
                v=up[h]
                # Browser copy uses ISO dates; Excel accepts these cleanly when pasted.
                arr[i]=v
            elif isinstance(arr[i],(datetime,date)):
                arr[i]=arr[i].strftime('%Y-%m-%d')
        rows.append(arr)
        if any(num(up.get(field))>0 for field in ('Paid Principal','Paid Profit','Late Profit')):
            ledger_updated_rows.append(arr)
    ledger_alloc={}
    for t in tx:
        if t.get('action')=='Investment Committed' and t.get('note'):
            ledger_alloc[t['note']]=ledger_alloc.get(t['note'],0)+num(t.get('amount'))
    from allocations import allocation_map, row_values, calculated_values
    all_saved_allocations=allocation_map()
    saved_allocations={code:record for code,record in all_saved_allocations.items()
                       if record.get('approval_status') in ('Approved','Disbursed')
                       or (include_drafts and record.get('approval_status') in ('Draft','Ready for Approval'))}
    added_entries=[];draft_preview_entries=[];future_entries=[];new_entry_rows=[]
    report_codes=set(template_codes)
    for code,record in sorted(saved_allocations.items()):
        if code in template_codes: continue
        approved=record.get('approval_status') in ('Approved','Disbursed')
        entry_date=parse_date(record.get('allocation_date')) or parse_date(record.get('disbursal_date'))
        if cutoff and entry_date and entry_date>cutoff:
            if approved:
                future_entries.append({'loan_code':code,'note_name':record.get('note_name',''),
                                       'investment_amount':record.get('investment_amount',0),
                                       'entry_date':entry_date.isoformat()})
            continue
        values=calculated_values(record)
        synthetic={
            'investment_amount':record.get('investment_amount',0),'final_date':values.get('Final Repayment Date '),
            'net_profit':values.get('Net Profit',0),'gross_profit':values.get('Gross Profit',0),
            'gross_profit_earned':values.get('Gross Profit Earned',0),'service_fee':values.get('Service Fee ',0),
            'late_profit':0,'loan_status':record.get('status','Active'),'installment':values.get('Installment',''),
            'early_repayment':'','early_date':None,
            'report_format':'expanded-profit' if 'Total Gross Profit' in headers else 'legacy'
        }
        # Reuse the same servicing calculation as historical rows by calculating
        # this entry against its ledger activity before it is appended.
        note_tx=[t for t in tx if t.get('note')==code]
        inv=max(0,num(synthetic['investment_amount']))
        paid_principal=sum(num(t.get('amount')) for t in note_tx if t.get('action')=='Principal Payout')
        profit_tx=[t for t in note_tx if t.get('action')=='Profit Payout']
        tawidh=sum(num(t.get('amount')) for t in profit_tx if 'tawidh' in str(t.get('transaction_no') or '').lower() or "ta'widh" in str(t.get('transaction_no') or '').lower())
        expanded=synthetic['report_format']=='expanded-profit'
        all_profit=sum(num(t.get('amount')) for t in profit_tx);paid_profit=all_profit if expanded else all_profit-tawidh
        from statements import profit_components
        sst=sum(float(profit_components(t.get('amount'),parse_date(t.get('date')))[2]) for t in profit_tx if parse_date(t.get('date')))
        principal_events=sorted((parse_date(t.get('date')),num(t.get('amount'))) for t in note_tx if t.get('action')=='Principal Payout' and parse_date(t.get('date')))
        running=0.0;settlement=None;tol=max(1.0,inv*.00001)
        for paid_on,amount in principal_events:
            running+=amount
            if inv and running>=inv-tol: settlement=paid_on;break
        completed=bool(inv and paid_principal>=inv-tol);final=parse_date(synthetic['final_date']);early=bool(completed and settlement and final and settlement<final)
        original_gross=max(0,num(synthetic['gross_profit']));original_net=max(0,num(synthetic['net_profit']))
        if completed and expanded:
            expected_net=paid_profit;gross_profit=(expected_net+sst)/.8 if expected_net or sst else 0
            gross_earned=gross_profit if early else max(0,num(synthetic['gross_profit_earned']))
            late_charge=0 if early else max(0,gross_profit-gross_earned);service_fee=max(0,gross_profit*.2);unpaid_profit=0
        elif completed:
            expected_net=paid_profit;gross_profit=expected_net/.8 if expected_net else original_gross
            gross_earned=max(0,num(synthetic['gross_profit_earned']));late_charge=0;service_fee=max(0,gross_profit-expected_net);unpaid_profit=0
        else:
            gross_profit=original_gross;gross_earned=max(0,num(synthetic['gross_profit_earned']));late_charge=0
            service_fee=max(0,num(synthetic['service_fee']));expected_net=max(0,gross_profit-service_fee-sst) if expanded else original_net
            unpaid_profit=max(0,expected_net-paid_profit)
        updates[code]={'_expanded':expanded,'Loan Status':'Completed' if completed else synthetic['loan_status'],
            'Actual Repayment **':paid_principal+paid_profit,'Paid Principal':paid_principal,'Unpaid Principal':max(0,inv-paid_principal),
            'Gross Profit':gross_profit,'Gross Profit Earned':gross_earned,'Total Gross Profit':gross_profit,
            'Net Profit':expected_net,'Paid Profit':paid_profit,'Unpaid Profit':unpaid_profit,'Late Profit':tawidh,
            'Late Payment Charges':late_charge,'Service Fee ':service_fee,'SST':sst,
            'Early Repayment':'Early Repayment of Principal' if early else '',
            'Early Repayment Date':settlement.isoformat() if early and settlement else None}
        arr=row_values(headers,record,len(rows)+1)
        for index,header in enumerate(headers):
            if header in updates[code]: arr[index]=updates[code][header]
        rows.append(arr);new_entry_rows.append(arr);template_codes.add(code)
        entry_summary={'loan_code':code,'note_name':record.get('note_name',''),
                       'issuer_name':record.get('issuer_name',''),'investment_amount':record.get('investment_amount',0),
                       'approval_status':record.get('approval_status','Draft'),
                       'paid_principal':paid_principal,'paid_profit':paid_profit}
        if approved:
            added_entries.append(entry_summary);report_codes.add(code)
        else: draft_preview_entries.append(entry_summary)
    missing=[{'loan_code':c,'allocated':a} for c,a in sorted(ledger_alloc.items()) if c not in report_codes]
    pending_entries=[{'loan_code':code,'note_name':record.get('note_name',''),
                      'investment_amount':record.get('investment_amount',0),
                      'approval_status':record.get('approval_status','Draft')}
                     for code,record in sorted(all_saved_allocations.items())
                     if record.get('approval_status') not in ('Approved','Disbursed','Cancelled')]
    positions={header:headers.index(header) for header in ('Loan Code','Issuer Name','Investment Amount','Paid Principal','Unpaid Principal','Paid Profit','Unpaid Profit','SST') if header in headers}
    def total(field):
        index=positions.get(field)
        return sum(num(row[index]) for row in rows if index is not None and index<len(row))
    report_totals={field:total(field) for field in ('Investment Amount','Paid Principal','Unpaid Principal','Paid Profit','Unpaid Profit','SST')}
    report_totals['new_allocation']=sum(num(item.get('investment_amount')) for item in added_entries)
    issues=[]
    if 'Loan Code' not in positions:
        issues.append({'level':'error','code':'missing-loan-code','message':'The report template does not contain a Loan Code column.'})
    else:
        code_counts={}
        for row in rows:
            code=normalize_note_code(row[positions['Loan Code']]) if positions['Loan Code']<len(row) else ''
            if code: code_counts[code]=code_counts.get(code,0)+1
        for code,count in code_counts.items():
            if count>1: issues.append({'level':'error','code':'duplicate-note','loan_code':code,'message':f'{code} appears {count} times in the report.'})
        investment_index=positions.get('Investment Amount');paid_principal_index=positions.get('Paid Principal')
        if investment_index is not None and paid_principal_index is not None:
            for row in rows:
                code=normalize_note_code(row[positions['Loan Code']]) if positions['Loan Code']<len(row) else ''
                investment=num(row[investment_index]) if investment_index<len(row) else 0
                paid_principal=num(row[paid_principal_index]) if paid_principal_index<len(row) else 0
                if code and paid_principal>investment+.01:
                    issues.append({'level':'error','code':'principal-overpayment','loan_code':code,
                                   'message':f'{code} principal payouts of RM {paid_principal:,.2f} exceed its RM {investment:,.2f} investment.'})
    for entry in added_entries:
        code=entry['loan_code'];investment=num(entry.get('investment_amount'));committed=num(ledger_alloc.get(code))
        if committed and abs(committed-investment)>.01:
            issues.append({'level':'warning','code':'ledger-allocation-mismatch','loan_code':code,
                           'message':f'{code} campaign allocation is RM {investment:,.2f}; the ledger records RM {committed:,.2f}.'})
        elif not committed:
            issues.append({'level':'warning','code':'ledger-allocation-missing','loan_code':code,
                           'message':f'{code} has no successful Investment Committed transaction through this cut-off.'})
    issuer_limit=num(get_settings().get('issuer_limit'));issuer_exposure={}
    issuer_index=positions.get('Issuer Name');outstanding_index=positions.get('Unpaid Principal');investment_index=positions.get('Investment Amount')
    if issuer_index is not None:
        for row in rows:
            issuer=str(row[issuer_index] if issuer_index<len(row) else '').strip() or 'Unknown Issuer'
            value=num(row[outstanding_index]) if outstanding_index is not None and outstanding_index<len(row) else (num(row[investment_index]) if investment_index is not None and investment_index<len(row) else 0)
            issuer_exposure[issuer]=issuer_exposure.get(issuer,0)+max(0,value)
        for issuer,exposure in issuer_exposure.items():
            if issuer_limit and exposure>issuer_limit+.01:
                issues.append({'level':'warning','code':'issuer-limit','issuer':issuer,
                               'message':f'{issuer} exposure is RM {exposure:,.2f}, above the RM {issuer_limit:,.2f} issuer limit.'})
    uncommitted=sum(num(entry.get('investment_amount')) for entry in added_entries if not num(ledger_alloc.get(entry['loan_code'])))
    latest_tx=max(tx,key=lambda row:(str(row.get('date') or ''),int(row.get('id') or 0)),default=None);available_cash=num(latest_tx.get('current_balance')) if latest_tx else 0
    if uncommitted>available_cash+.01:
        issues.append({'level':'warning','code':'cash-headroom',
                       'message':f'Uncommitted new allocations total RM {uncommitted:,.2f}, above the RM {available_cash:,.2f} ledger cash balance.'})
    if missing:
        issues.append({'level':'warning','code':'missing-campaign-details',
                       'message':f'{len(missing)} ledger allocation(s) still need approved campaign particulars.'})
    validation={'can_generate':not any(item['level']=='error' for item in issues),
                'error_count':sum(item['level']=='error' for item in issues),
                'warning_count':sum(item['level']=='warning' for item in issues),'issues':issues}
    result={'headers':headers,'rows':rows,'updates':updates,'missing_notes':missing,
            'baseline_row_count':len(displayed['rows']),'added_entries':added_entries,
            'draft_preview_entries':draft_preview_entries,'pending_entries':pending_entries,
            'future_entries':future_entries,'new_entry_rows':new_entry_rows,
            'ledger_updated_rows':ledger_updated_rows,'include_drafts':bool(include_drafts),
            'report_totals':report_totals,'validation':validation,
            'capacity':{'issuer_limit':issuer_limit,'issuer_exposure':issuer_exposure,
                        'available_cash':available_cash,'uncommitted_new_allocations':uncommitted},
            'as_of':as_of or (max([t.get('date') for t in tx if t.get('date')],default=''))}
    result['comparison']=None if include_drafts else simulation_comparison(headers,rows,compare_to)
    return result


def build_updated_simulation_workbook(as_of=None, snap=None):
    snap=snap or simulation_update_snapshot(as_of)
    wb=openpyxl.load_workbook(read_source('simulation'),data_only=False)
    ws=wb['Query result'] if 'Query result' in wb.sheetnames else wb[wb.sheetnames[0]]
    headers=[c.value for c in ws[1]]
    while headers and headers[-1] is None: headers.pop()
    hidx={h:i+1 for i,h in enumerate(headers)}
    for row in range(2,ws.max_row+1):
        code=normalize_note_code(ws.cell(row,2).value)
        up=snap['updates'].get(code)
        if not up: continue
        # Inputs sourced from the ledger. Derived cells retain the template's formula-driven style.
        input_fields=['Loan Status','Paid Principal','Paid Profit','Late Profit','Early Repayment','Early Repayment Date']
        if up.get('_expanded') and up.get('Loan Status')=='Completed':
            input_fields+=['Gross Profit Earned','Late Payment Charges','SST']
        for h in input_fields:
            if h not in hidx: continue
            v=up.get(h)
            if h=='Early Repayment Date' and v:
                v=parse_date(v)
            ws.cell(row,hidx[h]).value=v
        # When a note is completed early, its expected profit itself changes to the realised pro-rated amount.
        if up.get('Early Repayment') and up.get('Loan Status')=='Completed':
            if 'Gross Profit' in hidx: ws.cell(row,hidx['Gross Profit']).value=up['Gross Profit']
            if 'Net Profit' in hidx and 'Total Gross Profit' not in hidx: ws.cell(row,hidx['Net Profit']).value=up['Net Profit']
        # Re-apply the same core formulas used in the original MISB template for the transaction-driven fields.
        if 'Actual Repayment **' in hidx and 'Paid Principal' in hidx and 'Paid Profit' in hidx:
            ws.cell(row,hidx['Actual Repayment **']).value=f'={openpyxl.utils.get_column_letter(hidx["Paid Principal"])}{row}+{openpyxl.utils.get_column_letter(hidx["Paid Profit"])}{row}'
        if 'Unpaid Principal' in hidx and 'Investment Amount' in hidx and 'Paid Principal' in hidx:
            ws.cell(row,hidx['Unpaid Principal']).value=f'={openpyxl.utils.get_column_letter(hidx["Investment Amount"])}{row}-{openpyxl.utils.get_column_letter(hidx["Paid Principal"])}{row}'
        if 'Unpaid Profit' in hidx and 'Net Profit' in hidx and 'Paid Profit' in hidx:
            ws.cell(row,hidx['Unpaid Profit']).value=f'={openpyxl.utils.get_column_letter(hidx["Net Profit"])}{row}-{openpyxl.utils.get_column_letter(hidx["Paid Profit"])}{row}'
        if 'Service Fee ' in hidx and 'Total Gross Profit' in hidx:
            ws.cell(row,hidx['Service Fee ']).value=f'={openpyxl.utils.get_column_letter(hidx["Total Gross Profit"])}{row}*20%'
        elif 'Service Fee ' in hidx and 'Gross Profit' in hidx and 'Net Profit' in hidx:
            ws.cell(row,hidx['Service Fee ']).value=f'={openpyxl.utils.get_column_letter(hidx["Gross Profit"])}{row}-{openpyxl.utils.get_column_letter(hidx["Net Profit"])}{row}'
        if 'Net Profit' in hidx and 'Total Gross Profit' in hidx and 'Service Fee ' in hidx and 'SST' in hidx:
            ws.cell(row,hidx['Net Profit']).value=f'={openpyxl.utils.get_column_letter(hidx["Total Gross Profit"])}{row}-{openpyxl.utils.get_column_letter(hidx["Service Fee "])}{row}-{openpyxl.utils.get_column_letter(hidx["SST"])}{row}'
    from allocations import allocation_map, row_values
    existing={normalize_note_code(ws.cell(row,2).value) for row in range(2,ws.max_row+1)}
    included_codes={item['loan_code'] for item in snap.get('added_entries',[])}
    preview_rows={normalize_note_code(row[1]):row for row in snap.get('new_entry_rows',[]) if len(row)>1}
    for code,record in sorted(allocation_map(approved_only=True).items()):
        if code in existing or code not in included_codes: continue
        target=ws.max_row+1; source=max(2,target-1)
        for column in range(1,len(headers)+1):
            original=ws.cell(source,column); cell=ws.cell(target,column)
            if original.has_style:
                cell._style=copy(original._style)
            if original.number_format: cell.number_format=original.number_format
            if original.alignment: cell.alignment=copy(original.alignment)
            if original.protection: cell.protection=copy(original.protection)
        report_values=preview_rows.get(code) or row_values(headers,record,target-1)
        if report_values: report_values[0]=target-1
        for column,value in enumerate(report_values,1):
            ws.cell(target,column).value=value
        existing.add(code)
    # Keep every sheet, comment/note, style, column width and existing Remarks cell from the imported template.
    try:
        wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'
    except Exception:
        pass
    return wb,snap

@app.route('/api/simulation-preview')
def simulation_preview():
    as_of=request.args.get('as_of') or None
    include_drafts=request.args.get('include_drafts') in ('1','true','yes')
    compare_to=request.args.get('compare_to') or None
    if compare_to is not None:
        try: compare_to=int(compare_to)
        except (TypeError,ValueError): return jsonify({'error':'The comparison report ID is invalid.'}),400
    return jsonify(simulation_update_snapshot(as_of,include_drafts,compare_to))

@app.route('/api/export/updated-simulation.xlsx')
def export_updated_simulation():
    as_of=request.args.get('as_of') or None
    snap=simulation_update_snapshot(as_of)
    if not snap.get('validation',{}).get('can_generate',True):
        return jsonify({'error':'Resolve the blocking simulation validation errors before generating the report.',
                        'validation':snap['validation']}),422
    warning_count=int(snap.get('validation',{}).get('warning_count',0))
    warnings_reviewed=request.args.get('warnings_reviewed') in ('1','true','yes')
    if warning_count and not warnings_reviewed:
        return jsonify({'error':'Review and acknowledge the simulation warnings before generating the report.',
                        'validation':snap['validation']}),409
    wb,snap=build_updated_simulation_workbook(as_of,snap)
    b=io.BytesIO(); wb.save(b); b.seek(0)
    stamp=(as_of or snap.get('as_of') or date.today().isoformat());filename=f'Simulation_Report_MISB_Updated_{stamp}.xlsx'
    from storage import save_simulation_report_run
    save_simulation_report_run(stamp,filename,{
        'baseline_rows':snap.get('baseline_row_count',0),'new_rows':len(snap.get('added_entries',[])),
        'ledger_updated_rows':len(snap.get('ledger_updated_rows',[])),'total_rows':len(snap.get('rows',[])),
        'missing_allocations':len(snap.get('missing_notes',[]))
    },compact_simulation_snapshot(snap.get('headers',[]),snap.get('rows',[])))
    audit_event('Generated simulation report','report',stamp,{
        'filename':filename,'rows':len(snap.get('rows',[])),
        'warning_count':warning_count,'warnings_reviewed':warnings_reviewed})
    return send_file(b,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name=filename)


@app.get('/api/simulation-report-runs')
def simulation_report_history():
    from storage import simulation_report_runs
    rows=simulation_report_runs(request.args.get('limit',20))
    for row in rows: row.pop('snapshot',None)
    return jsonify(rows)

def report_workbook(p, as_of):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Fund Summary'
    green='0C7A58'; dark='17322B'; light='E8F4EF'; line='D9E4E0'
    ws['A1']='MISB Bi-Weekly Fund Monitoring Report'; ws['A1'].font=openpyxl.styles.Font(size=18,bold=True,color='FFFFFF'); ws['A1'].fill=openpyxl.styles.PatternFill('solid',fgColor=dark)
    ws.merge_cells('A1:D1'); ws['A2']='Report date'; ws['B2']=as_of; ws['A3']='Simulation source'; ws['B3']=p['summary']['simulation_as_of']; ws['A4']='Transaction ledger through'; ws['B4']=p['summary']['last_transaction_date']
    metrics=[('Cash available',p['summary']['cash']),('Total investment committed',p['summary']['total_investment_committed']),('Outstanding principal',p['summary']['outstanding_principal']),('Cumulative profit received',p['summary']['ledger_profit_paid']),('Return this quarter',p['summary']['return_this_quarter']),('Total returns since inception',p['summary']['total_return']),('Simulation outstanding principal',p['summary']['simulation_outstanding']),('Unpaid net profit',p['summary']['simulation_unpaid_profit']),('Issuer limit',p['settings']['issuer_limit'])]
    ws.append([]); ws.append(['Metric','Value'])
    for a,b in metrics: ws.append([a,b])
    for c in ws[6]: c.font=openpyxl.styles.Font(bold=True,color='FFFFFF'); c.fill=openpyxl.styles.PatternFill('solid',fgColor=green)
    for row in range(7,7+len(metrics)):
        ws.cell(row,2).number_format='0.00%' if 'return' in str(ws.cell(row,1).value).lower() else '"RM" #,##0.00'
    ws.column_dimensions['A'].width=36; ws.column_dimensions['B'].width=22; ws.column_dimensions['C'].width=18; ws.column_dimensions['D'].width=18

    sheets=[
      ('Issuer Exposure',['Issuer','Outstanding Principal','Limit','Remaining Allocation','Utilisation','Active Notes','Unpaid Profit'],
       [[x['issuer'],x['outstanding'],x['limit'],x['headroom'],x['utilization'],x['notes'],x['unpaid_profit']] for x in p['issuers']]),
      ('Receivables Calendar',['Date','Expected Principal','Notes Due'],[[x['date'],x['principal'],x['notes']] for x in p['receivables']]),
      ('Active Notes',['Loan Code','Issuer','Note Name','Status','Rating','Investment Amount','Outstanding Principal','Unpaid Profit','Final Repayment','Payment Type','Net p.m.','Monitoring Status','Owner','Next Action','Review Date'],
       [[r['loan_code'],r['issuer'],r['note_name'],r['loan_status'],r['rating'],r['investment_amount'],r['unpaid_principal'],r['unpaid_profit'],r['final_date'],r['payment_type'],r['net_pm'],r['monitoring'].get('status',''),r['monitoring'].get('owner',''),r['monitoring'].get('next_action',''),r['monitoring'].get('review_date','')] for r in p['portfolio'] if is_active(r)]),
      ('Transactions',['Date','Action','Note','Amount','Previous Balance','Current Balance','Transaction No','Pay Via'],
       [[r['date'],r['action'],r['note'],r['amount'],r['previous_balance'],r['current_balance'],r['transaction_no'],r['pay_via']] for r in p['transactions']])]
    for name,headers,rows in sheets:
        sh=wb.create_sheet(name); sh.append(headers)
        for c in sh[1]: c.font=openpyxl.styles.Font(bold=True,color='FFFFFF'); c.fill=openpyxl.styles.PatternFill('solid',fgColor=dark)
        for row in rows: sh.append(row)
        sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
        for col in range(1,len(headers)+1): sh.column_dimensions[openpyxl.utils.get_column_letter(col)].width=min(36,max(12,len(headers[col-1])+3))
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,(int,float)): c.number_format='#,##0.00'
    return wb

@app.route('/api/export/biweekly.xlsx')
def export_biweekly_xlsx():
    p=build_payload(); as_of=request.args.get('as_of') or date.today().isoformat(); wb=report_workbook(p,as_of)
    b=io.BytesIO(); wb.save(b); b.seek(0)
    return send_file(b,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name=f'MISB_Biweekly_Report_{as_of}.xlsx')


@app.route('/api/export/account-statement.xlsx')
def export_account_statement():
    month=request.args.get('month') or date.today().strftime('%Y-%m')
    try:
        start=datetime.strptime(month+'-01','%Y-%m-%d').date(); end=add_months(start,1)-timedelta(days=1)
    except:
        return jsonify({'error':'Month must be YYYY-MM'}),400
    rows=[]
    for r in load_transactions():
        d=parse_date(r.get('date'))
        if d and start<=d<=end and str(r.get('status','')).upper()=='SUCCESSFUL': rows.append(r)
    rows.sort(key=lambda r:(r.get('date') or '',r.get('id') or 0))
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Account Statement'
    dark='17322B'; green='0C7A58'
    ws['A1']='MISB Monthly Account Statement'; ws.merge_cells('A1:H1'); ws['A1'].font=openpyxl.styles.Font(size=18,bold=True,color='FFFFFF'); ws['A1'].fill=openpyxl.styles.PatternFill('solid',fgColor=dark)
    ws['A2']='Statement month'; ws['B2']=month
    opening=rows[0]['previous_balance'] if rows else 0; closing=rows[-1]['current_balance'] if rows else opening
    # Use the ledger's balance movement itself, so every action type is captured without assumptions.
    movements=[num(r.get('current_balance'))-num(r.get('previous_balance')) for r in rows]
    cash_in=sum(x for x in movements if x>0)
    cash_out=sum(-x for x in movements if x<0)
    ws['A4']='Opening balance'; ws['B4']=opening; ws['A5']='Cash in'; ws['B5']=cash_in; ws['A6']='Cash out'; ws['B6']=cash_out; ws['A7']='Closing balance'; ws['B7']=closing
    for rr in range(4,8): ws.cell(rr,2).number_format='"RM" #,##0.00'
    headers=['Date','Action','Note','Transaction No','Amount (RM)','Previous Balance (RM)','Current Balance (RM)','Pay Via']
    ws.append([]); ws.append(headers)
    hr=ws.max_row
    for c in ws[hr]: c.font=openpyxl.styles.Font(bold=True,color='FFFFFF'); c.fill=openpyxl.styles.PatternFill('solid',fgColor=green)
    for r in rows:
        ws.append([r['date'],r['action'],r['note'],r['transaction_no'],r['amount'],r['previous_balance'],r['current_balance'],r['pay_via']])
    for row in ws.iter_rows(min_row=hr+1):
        for c in row[4:7]: c.number_format='"RM" #,##0.00'
    widths=[14,24,14,36,18,22,22,14]
    for i,w in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes=f'A{hr+1}'; ws.auto_filter.ref=f'A{hr}:H{ws.max_row}'
    b=io.BytesIO(); wb.save(b); b.seek(0)
    return send_file(b,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name=f'MISB_Account_Statement_{month}.xlsx')

@app.route('/api/export/biweekly.csv')
def export_biweekly():
    p=build_payload(); output=io.StringIO(); w=csv.writer(output); s=p['summary']
    w.writerow(['MISB Bi-Weekly Monitoring Snapshot', datetime.now().strftime('%Y-%m-%d')]); w.writerow([]); w.writerow(['Metric','Amount / Count'])
    for k in ['cash','ledger_deployed','implied_fund_value','ledger_profit_paid','simulation_outstanding','simulation_unpaid_profit','active_notes','portfolio_gap']:
        w.writerow([k,s[k]])
    w.writerow([]); w.writerow(['Issuer','Outstanding','Limit','Remaining Allocation','Utilisation'])
    for x in p['issuers']: w.writerow([x['issuer'],x['outstanding'],x['limit'],x['headroom'],x['utilization']])
    b=io.BytesIO(output.getvalue().encode('utf-8-sig')); b.seek(0)
    return send_file(b,mimetype='text/csv',as_attachment=True,download_name='MISB_Biweekly_Monitoring.csv')


from web import configure_web
configure_web(app)
from statements import register_statement_routes
register_statement_routes(app)
from allocations import register_allocation_routes
register_allocation_routes(app)
