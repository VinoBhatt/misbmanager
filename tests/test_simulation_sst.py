import io
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT=Path(__file__).resolve().parent.parent
sys.path[:0]=[str(ROOT/'src'),str(ROOT/'scripts')]
from app import app, simulation_update_snapshot, build_updated_simulation_workbook
from local import initialize
from storage import save_source


class SimulationSSTTests(unittest.TestCase):
    def test_sst_only_uses_successful_profit_payouts_from_august_first(self):
        with tempfile.TemporaryDirectory() as directory:
            initialize(Path(directory))
            with patch.dict(os.environ,{'MISB_DATA_DIR':directory}):
                values={'No.':1,'Loan Code':'IIF-1234','Issuer Name':'Example issuer',
                        'Investment Amount':1000,'Loan Status':'Active','Gross Profit Earned':200,
                        'Total Gross Profit':200,'Late Payment Charges':0,'Service Fee ':40,
                        'SST':99,'Net Profit':61,'Paid Principal':0,'Unpaid Principal':1000,
                        'Paid Profit':0,'Unpaid Profit':61,'Disbursal Date':'2026-07-01',
                        'Final Repayment Date ':'2026-12-01'}
                workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='Query result'
                sheet.append(list(values));sheet.append(list(values.values()))
                output=io.BytesIO();workbook.save(output)
                with app.test_request_context():
                    save_source('simulation',output.getvalue(),'2026-09-02')
                def transaction(action,amount,day,status='SUCCESSFUL'):
                    return {'note':'IIF-1234','action':action,'amount':amount,'date':day,
                            'status':status,'transaction_no':''}
                july=transaction('Profit Payout',80,'2026-07-31')
                august=transaction('Profit Payout',78.4,'2026-08-01')
                cases=[
                    ('unpaid profit',[],None,0,0),
                    ('before August',[july],None,0,80),
                    ('August boundary',[august],None,1.6,78.4),
                    ('pending payout',[{**august,'status':'PENDING'}],None,0,0),
                    ('failed payout',[{**august,'status':'FAILED'}],None,0,0),
                    ('principal only',[transaction('Principal Payout',100,'2026-08-01')],None,0,0),
                    ('mixed payout dates',[july,august],None,1.6,158.4),
                    ('report before payout',[july,august],'2026-07-31',0,80),
                ]
                for label,payouts,cutoff,expected_sst,expected_paid in cases:
                    with self.subTest(label=label),app.test_request_context(),patch('app.load_transactions',
                            return_value=[transaction('Investment Committed',1000,'2026-07-01'),*payouts]):
                        snapshot=simulation_update_snapshot(cutoff)
                        row=dict(zip(snapshot['headers'],snapshot['rows'][0]))
                        self.assertAlmostEqual(row['SST'],expected_sst)
                        self.assertAlmostEqual(row['Paid Profit'],expected_paid)
                        self.assertAlmostEqual(row['Net Profit'],160-expected_sst)
                        self.assertAlmostEqual(row['Unpaid Profit'],160-expected_sst-expected_paid)
                        exported,_=build_updated_simulation_workbook(snap=snapshot)
                        self.assertAlmostEqual(exported['Query result'].cell(2,list(values).index('SST')+1).value,expected_sst)
                        exported.close()

    def test_example_matches_preview_and_excel_for_active_and_completed_notes(self):
        self.check_example(expanded=True)

    def test_legacy_report_keeps_platform_fee_separate_from_sst(self):
        self.check_example(expanded=False)

    def test_prior_net_amount_does_not_cause_sst_to_be_deducted_twice(self):
        self.check_example(expanded=False,recorded_net=78.4)

    def check_example(self,expanded,recorded_net=80):
        with tempfile.TemporaryDirectory() as directory:
            initialize(Path(directory))
            with patch.dict(os.environ,{'MISB_DATA_DIR':directory}):
                values={'No.':1,'Loan Code':'IIF-1234','Issuer Name':'Example issuer',
                        'Investment Amount':1000,'Loan Status':'Active','Gross Profit Earned':100,
                        'Total Gross Profit':100,'Late Payment Charges':0,'Service Fee ':20,'SST':0,
                        'Net Profit':recorded_net,'Paid Principal':0,'Unpaid Principal':1000,'Paid Profit':0,
                        'Unpaid Profit':80,'Final Repayment Date ':'2026-09-02'}
                if not expanded:
                    for field in ('Gross Profit Earned','Total Gross Profit','Late Payment Charges','SST'):
                        values.pop(field)
                    values['Gross Profit']=100
                workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='Query result'
                sheet.append(list(values));sheet.append(list(values.values()))
                output=io.BytesIO();workbook.save(output)
                with app.test_request_context():
                    save_source('simulation',output.getvalue(),'2026-09-02')
                base={'note':'IIF-1234','status':'SUCCESSFUL','date':'2026-09-02','transaction_no':''}
                for completed in (False,True):
                    with self.subTest(completed=completed):
                        transactions=[{**base,'action':'Investment Committed','amount':1000},
                                      {**base,'action':'Profit Payout','amount':78.40}]
                        if completed: transactions.append({**base,'action':'Principal Payout','amount':1000})
                        with app.test_request_context(),patch('app.load_transactions',return_value=transactions):
                            snapshot=simulation_update_snapshot()
                            row=dict(zip(snapshot['headers'],snapshot['rows'][0]))
                            self.assertEqual(row['Loan Status'],'Completed' if completed else 'Active')
                            gross_field='Total Gross Profit' if expanded else 'Gross Profit'
                            self.assertAlmostEqual(snapshot['updates']['IIF-1234']['SST'],1.6)
                            for field,expected in [(gross_field,100),('Service Fee ',20),('Net Profit',78.4),('Paid Profit',78.4),('Unpaid Profit',0)]:
                                self.assertAlmostEqual(row[field],expected,msg=field)
                            exported,_=build_updated_simulation_workbook(snap=snapshot)
                            saved=io.BytesIO();exported.save(saved);saved.seek(0)
                            reloaded=openpyxl.load_workbook(saved)
                            report=reloaded['Query result']
                            cell=lambda name: report.cell(2,list(values).index(name)+1)
                            gross_column=openpyxl.utils.get_column_letter(list(values).index(gross_field)+1)
                            self.assertEqual(cell('Service Fee ').value,f'={gross_column}2*20%')
                            if expanded:
                                self.assertEqual(cell('SST').value,1.6)
                                self.assertEqual(cell('Net Profit').value,'=G2-I2-J2')
                            else:
                                self.assertAlmostEqual(cell('Net Profit').value,78.4)
                                self.assertAlmostEqual(cell('Gross Profit').value,100)
                            reloaded.close();exported.close()


if __name__=='__main__':unittest.main()
