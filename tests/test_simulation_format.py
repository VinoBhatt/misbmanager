import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

ROOT=Path(__file__).resolve().parent.parent
sys.path[:0]=[str(ROOT/'src'),str(ROOT/'scripts')]
from app import app, load_simulation, simulation_update_snapshot, build_updated_simulation_workbook
from local import initialize
from simulation_format import realised_values, NEW_REQUIRED
from storage import save_source
from uploads import inspect_workbook


class SeptemberFormatTests(unittest.TestCase):
    def test_actual_components_and_early_settlement_preserve_projection(self):
        payouts=[{'amount':80,'date':'2026-07-31'},
                 {'amount':78.4,'date':'2026-08-01'},
                 {'amount':7.84,'date':'2026-09-01','transaction_no':'Tawidh'}]
        active=realised_values(500,payouts)
        self.assertAlmostEqual(active['Total Gross Profit Earned'],210)
        self.assertAlmostEqual(active['Total Service Fee Charged'],42)
        self.assertAlmostEqual(active['SST Charged'],1.76)
        self.assertAlmostEqual(active['Unpaid Profit'],233.76)
        self.assertAlmostEqual(active['Net Late Payment Charges (SST Affected)'],7.84)
        early=realised_values(500,payouts,True)
        self.assertEqual(early['Total Net Profit Projected'],400)
        self.assertEqual(early['Unpaid Profit'],0)
        self.assertEqual(early['SST Control Check'],.08)

    def test_import_preview_and_export_use_new_headers(self):
        values={'No.':1,'Loan Code':'IIF-1234','Issuer Name':'Example issuer',
                'Investment Amount':1000,'Loan Status':'Active',
                'Paid Principal':0,'Unpaid Principal':1000,'Unpaid Profit':400,
                'Gross Profit Projected':500,'Total Net Profit Projected':400,
                'Final Repayment Date ':'2026-12-31','Early Repayment':'',
                'Early Repayment Date':'','Remarks':'Preserve this remark'}
        for field in sorted(NEW_REQUIRED): values.setdefault(field,0)
        workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='Query result'
        sheet.append(list(values));sheet.append(list(values.values()))
        output=io.BytesIO();workbook.save(output)
        self.assertTrue(inspect_workbook('simulation',output.getvalue())['can_import'])
        with tempfile.TemporaryDirectory() as directory:
            initialize(Path(directory))
            with patch.dict(os.environ,{'MISB_DATA_DIR':directory}),app.test_request_context():
                save_source('simulation',output.getvalue(),'2026-09-30')
                portfolio,_=load_simulation()
                self.assertEqual(portfolio[0]['report_format'],'projected-realised')
                self.assertEqual(portfolio[0]['net_profit'],400)
                transactions=[{'note':'IIF-1234','action':'Profit Payout','amount':78.4,
                               'date':'2026-09-01','status':'SUCCESSFUL'}]
                with patch('app.load_transactions',return_value=transactions):
                    snapshot=simulation_update_snapshot('2026-09-30')
                row=dict(zip(snapshot['headers'],snapshot['rows'][0]))
                self.assertAlmostEqual(row['Total Gross Profit Earned'],100)
                self.assertAlmostEqual(row['Total Paid Out Profit'],78.4)
                self.assertAlmostEqual(row['Unpaid Profit'],321.6)
                self.assertEqual(snapshot['report_totals']['Paid Profit'],78.4)
                exported,_=build_updated_simulation_workbook(snap=snapshot)
                sheet=exported['Query result'];columns={c.value:c.column for c in sheet[1]}
                self.assertEqual(sheet.cell(2,columns['Gross Profit Projected']).value,500)
                self.assertEqual(sheet.cell(2,columns['Total Net Profit Projected']).value,400)
                self.assertEqual(sheet.cell(2,columns['SST Net Profit (After 1 Aug)']).value,78.4)
                self.assertIn('MAX(0,',sheet.cell(2,columns['Unpaid Profit']).value)
                self.assertEqual(sheet.cell(2,columns['Remarks']).value,'Preserve this remark')
                exported.close()


if __name__=='__main__': unittest.main()
