from contextlib import closing
import importlib.util
import io
import os
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch
import openpyxl

ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT/'src'))
sys.path.insert(0,str(ROOT/'scripts'))
from local import initialize
from app import app


class WebsiteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.root=Path(cls.temp.name)
        with zipfile.ZipFile(ROOT/'MISB_Fund_Tracker_simulation_copy_FIXED_v2.zip') as archive:
            archive.extractall(cls.root/'original')
        original=cls.root/'original/MISB_Fund_Tracker'
        spec=importlib.util.spec_from_file_location('original_app',original/'app.py')
        cls.original=importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.original)
        cls.original.init_db()
        cls.original_payload=cls.original.build_payload()
        cls.seed=original/'data'

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def setUp(self):
        self.directory=tempfile.TemporaryDirectory()
        initialize(Path(self.directory.name),self.seed)
        self.env=patch.dict(os.environ,{'MISB_DATA_DIR':self.directory.name,'MISB_LOCAL_PREVIEW':'1'})
        self.env.start()
        self.client=app.test_client()

    def tearDown(self):
        self.env.stop()
        self.directory.cleanup()

    def test_original_dashboard_parity(self):
        response=self.client.get('/api/data')
        self.assertEqual(response.status_code,200)
        for key in ('summary','portfolio','issuers','profit_schedule','cash_projection','transactions'):
            # JSON round-trip normalizes date values exactly as both HTTP APIs do.
            with self.original.app.app_context():
                expected=self.original.app.json.loads(self.original.app.json.dumps(self.original_payload[key]))
            self.assertEqual(response.json[key],expected,key)

    def test_saved_changes_survive_new_client(self):
        self.assertEqual(self.client.post('/api/settings',json={'issuer_limit':1500000}).status_code,200)
        self.assertEqual(self.client.post('/api/monitoring/IIF-12345',json={'status':'Watch','owner':'Manager','notes':'Review'}).status_code,200)
        response=self.client.post('/api/cashflow-plan',json={'flow_type':'Allocation','flow_date':'2027-01-01','amount':1200})
        self.assertEqual(response.status_code,200)
        data=app.test_client().get('/api/data').json
        self.assertEqual(data['settings']['issuer_limit'],1500000)
        self.assertTrue(any(r['id']==response.json['id'] for r in data['cashflow_plan']))
        self.assertEqual(self.client.delete('/api/cashflow-plan/'+str(response.json['id'])).status_code,200)
        portfolio=data['portfolio']
        note=next(r for r in portfolio if any(not e['is_paid'] for e in r['payment_schedule']))
        event=next(e for e in note['payment_schedule'] if not e['is_paid'])
        mark={**event,'loan_code':note['loan_code'],'is_paid':True,'paid_date':'2026-08-24','paid_amount':123.45}
        self.assertEqual(self.client.post('/api/payment-mark',json=mark).status_code,200)
        refreshed=app.test_client().get('/api/data').json
        saved=next(r for r in refreshed['portfolio'] if r['loan_code']==note['loan_code'])
        payment=next(r for r in saved['payment_schedule'] if r['event_key']==event['event_key'])
        self.assertEqual(payment['paid_amount'],123.45)
        self.assertEqual(payment['paid_date'],'2026-08-24')

    def test_all_excel_exports(self):
        for endpoint in ('biweekly.xlsx','account-statement.xlsx?month=2026-08','updated-simulation.xlsx?as_of=2026-08-15'):
            response=self.client.get('/api/export/'+endpoint)
            self.assertEqual(response.status_code,200,endpoint)
            wb=openpyxl.load_workbook(io.BytesIO(response.data))
            self.assertTrue(wb.sheetnames)
            if endpoint=='biweekly.xlsx':
                cash=next(row[1].value for row in wb['Fund Summary'].iter_rows() if row[0].value=='Cash available')
                self.assertEqual(cash,self.original_payload['summary']['cash'])
            if endpoint.startswith('updated'):
                original=openpyxl.load_workbook(self.seed/'simulation.xlsx')
                self.assertEqual(wb.sheetnames,original.sheetnames)
                sheet=wb['Query result']; old=original['Query result']
                headers=[c.value for c in sheet[1]]
                col=headers.index('Remarks')+1
                self.assertEqual([sheet.cell(r,col).value for r in range(2,sheet.max_row+1)],
                                 [old.cell(r,col).value for r in range(2,old.max_row+1)])
                original.close()
            wb.close()

    def test_simulation_report_history_and_comparison(self):
        before=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-21'}).json
        self.assertFalse(before['comparison']['has_prior'])
        generated=self.client.get('/api/export/updated-simulation.xlsx',query_string={'as_of':'2026-09-21','warnings_reviewed':'1'})
        self.assertEqual(generated.status_code,200)
        history=self.client.get('/api/simulation-report-runs').json
        self.assertEqual(len(history),1)
        self.assertEqual(history[0]['as_of'],'2026-09-21')
        self.assertEqual(history[0]['total_rows'],len(before['rows']))
        self.assertNotIn('snapshot',history[0])
        selected=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-21','compare_to':history[0]['id']}).json
        self.assertTrue(selected['comparison']['has_prior'])
        self.assertEqual(selected['comparison']['prior_id'],history[0]['id'])
        missing=self.client.get('/api/simulation-preview',query_string={'compare_to':999999}).json
        self.assertFalse(missing['comparison']['has_prior'])
        self.assertEqual(missing['comparison']['requested_id'],999999)
        self.assertEqual(self.client.get('/api/simulation-preview',query_string={'compare_to':'bad'}).status_code,400)
        after=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-21'}).json
        self.assertTrue(after['comparison']['has_prior'])
        self.assertEqual(after['comparison']['new_notes'],[])
        self.assertEqual(after['comparison']['removed_notes'],[])
        self.assertTrue(any(row['action']=='Generated simulation report' for row in self.client.get('/api/audit-events').json))

    def test_new_note_allocation_feeds_simulation_copy(self):
        snapshot=self.client.get('/api/simulation-preview').json
        code_index=snapshot['headers'].index('Loan Code');investment_index=snapshot['headers'].index('Investment Amount')
        baseline_row=next(row for row in snapshot['rows'] if row[code_index] and float(row[investment_index] or 0)>0)
        baseline_overpayment={'id':999997,'date':'2026-09-20','action':'Principal Payout',
                              'amount':float(baseline_row[investment_index])+1,'note':baseline_row[code_index],
                              'status':'SUCCESSFUL','transaction_no':'baseline overpayment test',
                              'current_balance':0,'previous_balance':0,'pay_via':''}
        with patch('app.load_transactions',return_value=self.client.get('/api/data').json['transactions']+[baseline_overpayment]):
            baseline_invalid=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-21'}).json
        self.assertTrue(any(issue['code']=='principal-overpayment' and issue['loan_code']==baseline_row[code_index]
                            for issue in baseline_invalid['validation']['issues']))
        missing=snapshot['missing_notes'][0]
        self.assertEqual(snapshot['baseline_row_count'],len(snapshot['rows']))
        self.assertEqual(snapshot['added_entries'],[])
        pending=self.client.get('/api/data').json['pending_allocations']
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in pending])
        self.assertEqual(next(item['allocated'] for item in pending if item['loan_code']==missing['loan_code']),missing['allocated'])
        payload={
            'loan_code':missing['loan_code'],'company_id':6001,'issuer_name':'Example Manufacturing Sdn Bhd',
            'note_name':'Precision Tools Manufacturer 11','product_type':'Islamic Invoice Financing (IIF) - Receivables',
            'rating':'CR6','business_description':'Precision component manufacturing and invoice financing.',
            'investment_amount':missing['allocated'],'loan_note_size':386000,'allocation_date':'2026-09-17',
            'disbursal_date':'2026-09-25','payment_type':'Profit Only','term':90,'tenor_type':'Days',
            'gross_pa':15.6,'campaign_start':'2026-09-17','campaign_end':'2026-09-24',
            'status':'Active','remarks':'Generated from note card'
        }
        duplicate={**payload,'loan_code':self.client.get('/api/data').json['portfolio'][0]['loan_code']}
        response=self.client.post('/api/note-allocations',json=duplicate)
        self.assertEqual(response.status_code,409,response.data)
        self.assertIn('historical simulation baseline',response.json['error'])
        response=self.client.post('/api/note-allocations',json=payload)
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.json['gross_pa'],.156)
        draft=self.client.get('/api/simulation-preview').json
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in draft['missing_notes']])
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in draft['pending_entries']])
        self.assertEqual(len(draft['rows']),draft['baseline_row_count'])
        draft_preview=self.client.get('/api/simulation-preview',query_string={'include_drafts':'1'}).json
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in draft_preview['draft_preview_entries']])
        self.assertIn(missing['loan_code'],[row[draft_preview['headers'].index('Loan Code')] for row in draft_preview['new_entry_rows']])
        self.assertEqual(len(draft_preview['rows']),draft_preview['baseline_row_count']+1)
        # A draft preview never changes the final downloadable workbook.
        draft_export=self.client.get('/api/export/updated-simulation.xlsx',query_string={'warnings_reviewed':'1'})
        draft_workbook=openpyxl.load_workbook(io.BytesIO(draft_export.data),data_only=True)
        self.assertNotIn(missing['loan_code'],[draft_workbook['Query result'].cell(row,2).value for row in range(2,draft_workbook['Query result'].max_row+1)])
        draft_workbook.close()
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in self.client.get('/api/data').json['pending_allocations']])
        response=self.client.patch('/api/note-allocations/'+missing['loan_code']+'/approval',json={'status':'Ready for Approval'})
        self.assertEqual(response.status_code,200,response.data)
        response=self.client.patch('/api/note-allocations/'+missing['loan_code']+'/approval',json={'status':'Approved'})
        self.assertEqual(response.status_code,200,response.data)
        self.assertNotIn(missing['loan_code'],[item['loan_code'] for item in self.client.get('/api/data').json['pending_allocations']])
        before_entry=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-16'}).json
        self.assertNotIn(missing['loan_code'],[item['loan_code'] for item in before_entry['added_entries']])
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in before_entry['future_entries']])
        edited={**payload,'remarks':'Reviewed campaign particulars'}
        self.assertEqual(self.client.post('/api/note-allocations',json=edited).status_code,200)
        saved=next(row for row in self.client.get('/api/note-allocations').json if row['loan_code']==missing['loan_code'])
        self.assertEqual(saved['approval_status'],'Approved')
        history=self.client.get('/api/note-allocations/'+missing['loan_code']+'/history').json
        self.assertEqual([row['action'] for row in history],['Updated','Approval changed','Approval changed','Created'])
        self.assertEqual(history[0]['changed_fields'],['remarks'])
        transactions=self.client.get('/api/data').json['transactions']
        repayment={'id':999999,'date':'2026-09-20','action':'Principal Payout','amount':1000,
                   'note':missing['loan_code'],'status':'SUCCESSFUL','transaction_no':'test repayment',
                   'current_balance':0,'previous_balance':0,'pay_via':''}
        with patch('app.load_transactions',return_value=transactions+[repayment]):
            serviced=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-21'}).json
        loan_index=serviced['headers'].index('Loan Code');paid_index=serviced['headers'].index('Paid Principal')
        generated=next(row for row in serviced['new_entry_rows'] if row[loan_index]==missing['loan_code'])
        self.assertEqual(generated[paid_index],1000)
        self.assertTrue(serviced['validation']['can_generate'])
        self.assertGreater(serviced['report_totals']['Investment Amount'],0)
        self.assertEqual(serviced['capacity']['issuer_limit'],self.client.get('/api/data').json['settings']['issuer_limit'])
        self.assertIn('available_cash',serviced['capacity'])
        overpayment={**repayment,'id':999998,'amount':missing['allocated']+1}
        with patch('app.load_transactions',return_value=transactions+[overpayment]):
            invalid=self.client.get('/api/simulation-preview',query_string={'as_of':'2026-09-21'}).json
            blocked=self.client.get('/api/export/updated-simulation.xlsx',query_string={'as_of':'2026-09-21'})
        self.assertFalse(invalid['validation']['can_generate'])
        self.assertTrue(any(issue['code']=='principal-overpayment' for issue in invalid['validation']['issues']))
        self.assertEqual(blocked.status_code,422)
        email=self.client.get('/api/allocation-email-preview',query_string={
            'loan_code':missing['loan_code'],'request_date':'2026-09-17','period_end':'2026-09-30',
            'additional_amount':'7000000','additional_date':'2026-09-01',
            'reserve_amount':'645175.47','reserve_date':'2026-09-02'
        })
        self.assertEqual(email.status_code,200,email.data)
        self.assertEqual(email.json['reference_number'],missing['loan_code'].replace('-','')+'-17092026')
        self.assertAlmostEqual(email.json['total_expected'],email.json['available']+email.json['expected']-email.json['reserve'])
        self.assertAlmostEqual(email.json['net_available'],email.json['total_expected']-missing['allocated'])
        self.assertIn('Dear Muamalat Invest Operations Team',email.json['plain_text'])
        automatic=self.client.get('/api/allocation-email-preview',query_string={
            'loan_code':missing['loan_code'],'request_date':'2026-09-17','period_end':'2026-09-30'
        })
        self.assertEqual(automatic.status_code,200,automatic.data)
        eligible=[row for row in self.client.get('/api/data').json['transactions']
                  if row['status']=='SUCCESSFUL' and row['action']=='Deposit Approved' and row['date']<='2026-09-17']
        if eligible:
            latest=max(row['date'] for row in eligible)
            self.assertEqual(automatic.json['additional_date'],latest)
            self.assertEqual(automatic.json['additional_amount'],sum(row['amount'] for row in eligible if row['date']==latest))
        refreshed=self.client.get('/api/simulation-preview').json
        self.assertNotIn(missing['loan_code'],[item['loan_code'] for item in refreshed['missing_notes']])
        self.assertIn(missing['loan_code'],[item['loan_code'] for item in refreshed['added_entries']])
        self.assertEqual(len(refreshed['rows']),refreshed['baseline_row_count']+1)
        index=refreshed['headers'].index('Loan Code')
        self.assertIn(missing['loan_code'],[row[index] for row in refreshed['rows']])
        warning_gate=self.client.get('/api/export/updated-simulation.xlsx')
        self.assertEqual(warning_gate.status_code,409)
        self.assertIn('acknowledge',warning_gate.json['error'])
        exported=self.client.get('/api/export/updated-simulation.xlsx',query_string={'warnings_reviewed':'1'})
        workbook=openpyxl.load_workbook(io.BytesIO(exported.data),data_only=True)
        sheet=workbook['Query result']
        self.assertIn(missing['loan_code'],[sheet.cell(row,2).value for row in range(2,sheet.max_row+1)])
        workbook.close()
        material={**edited,'investment_amount':missing['allocated']-1}
        changed=self.client.post('/api/note-allocations',json=material)
        self.assertEqual(changed.status_code,200,changed.data)
        self.assertTrue(changed.json['approval_reset'])
        self.assertIn('investment_amount',changed.json['changed_fields'])
        saved=next(row for row in self.client.get('/api/note-allocations').json if row['loan_code']==missing['loan_code'])
        self.assertEqual(saved['approval_status'],'Draft')
        versions=self.client.get('/api/note-allocations/'+missing['loan_code']+'/history').json
        self.assertEqual(len(versions),5)
        restored=self.client.post(f"/api/note-allocations/{missing['loan_code']}/history/{versions[1]['id']}/restore")
        self.assertEqual(restored.status_code,200,restored.data)
        self.assertEqual(restored.json['approval_status'],'Draft')
        self.assertEqual(restored.json['investment_amount'],missing['allocated'])
        versions=self.client.get('/api/note-allocations/'+missing['loan_code']+'/history').json
        self.assertEqual(versions[0]['action'],'Restored')
        self.assertEqual(len(versions),6)
        self.assertEqual(self.client.delete('/api/note-allocations/'+missing['loan_code']).status_code,200)

    def test_invalid_upload_does_not_replace_source(self):
        before=self.client.get('/api/data').json['summary']
        response=self.client.post('/api/upload/transactions',data={'file':(io.BytesIO(b'not excel'),'bad.xlsx')})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.client.get('/api/data').json['summary'],before)

    def test_valid_upload_preserves_backup(self):
        response=self.client.post('/api/upload/transactions',data={'file':(io.BytesIO((self.seed/'transactions.xlsx').read_bytes()),'transactions.xlsx')})
        self.assertEqual(response.status_code,200)
        import sqlite3
        with closing(sqlite3.connect(Path(self.directory.name)/'misb_tracker.db')) as con, con:
            self.assertIsNotNone(con.execute('SELECT 1 FROM workbook_versions WHERE object_key=?',(response.json['backup'],)).fetchone())
        self.assertEqual(self.client.get('/api/data').json['summary'],self.original_payload['summary'])
        versions=self.client.get('/api/source-versions/transactions').json
        self.assertTrue(any(row['current'] for row in versions))
        current=next(row for row in versions if row['current'])
        self.assertEqual(current['filename'],'transactions.xlsx')
        old=next(row for row in versions if row['object_key']==response.json['backup'])
        before_key=current['object_key']
        downloaded=self.client.get('/api/source-versions/transactions/download',
                                   query_string={'object_key':old['object_key']})
        self.assertEqual(downloaded.status_code,200,downloaded.data[:100])
        self.assertEqual(downloaded.data,(self.seed/'transactions.xlsx').read_bytes())
        self.assertEqual(next(row for row in self.client.get('/api/source-versions/transactions').json
                              if row['current'])['object_key'],before_key)
        self.assertEqual(self.client.get('/api/source-versions/simulation/download',
            query_string={'object_key':old['object_key']}).status_code,404)
        self.assertEqual(self.client.get('/api/source-versions/transactions/download',
            query_string={'object_key':'sources/transactions/missing.xlsx'}).status_code,404)
        restored=self.client.post('/api/source-versions/transactions/restore',json={'object_key':old['object_key']})
        self.assertEqual(restored.status_code,200,restored.data)
        self.assertTrue(next(row for row in self.client.get('/api/source-versions/transactions').json
                             if row['object_key']==old['object_key'])['current'])

    def test_validation_report_does_not_replace_live_source(self):
        before=self.client.get('/api/data').json
        response=self.client.post('/api/upload/validate/transactions',
            data={'file':(io.BytesIO((self.seed/'transactions.xlsx').read_bytes()),'transactions.xlsx')})
        self.assertEqual(response.status_code,200,response.data)
        self.assertTrue(response.json['can_import'])
        self.assertGreater(response.json['stats']['successful'],0)
        self.assertEqual(response.json['stats']['balance_breaks'],0)
        self.assertEqual(response.json['stats']['movement_mismatches'],0)
        self.assertGreater(response.json['stats']['excluded_cash_shadows'],0)
        self.assertEqual(response.json['comparison']['added'],0)
        self.assertEqual(response.json['comparison']['changed'],0)
        self.assertEqual(response.json['comparison']['removed'],0)
        self.assertEqual(response.json['comparison']['current_balance'],response.json['comparison']['candidate_balance'])
        after=self.client.get('/api/data').json
        self.assertEqual(after['summary'],before['summary'])

        workbook=openpyxl.load_workbook(self.seed/'transactions.xlsx')
        sheet=workbook[workbook.sheetnames[0]];headers=[cell.value for cell in sheet[2]]
        current_column=headers.index('Current Balance(MYR)')+1
        sheet.cell(4,current_column).value=float(sheet.cell(4,current_column).value)+10
        broken=io.BytesIO();workbook.save(broken);workbook.close();payload=broken.getvalue()
        rejected=self.client.post('/api/upload/validate/transactions',
            data={'file':(io.BytesIO(payload),'broken-balances.xlsx')})
        self.assertEqual(rejected.status_code,200,rejected.data)
        self.assertFalse(rejected.json['can_import'])
        messages=[item['message'] for item in rejected.json['issues']]
        self.assertIn('Ledger balance is not continuous between transactions',messages)
        self.assertIn('Transaction amount does not match its balance movement',messages)
        self.assertGreaterEqual(rejected.json['comparison']['changed'],1)
        direct=self.client.post('/api/upload/transactions',
            data={'file':(io.BytesIO(payload),'broken-balances.xlsx')})
        self.assertEqual(direct.status_code,422,direct.data)
        self.assertEqual(self.client.get('/api/data').json['summary'],before['summary'])

    def test_source_regression_requires_explicit_acknowledgement(self):
        workbook=openpyxl.load_workbook(self.seed/'transactions.xlsx');sheet=workbook[workbook.sheetnames[0]]
        headers=[cell.value for cell in sheet[2]];removed_id=str(sheet.cell(3,headers.index('ID')+1).value);sheet.delete_rows(3)
        stream=io.BytesIO();workbook.save(stream);workbook.close();candidate=stream.getvalue()
        preview=self.client.post('/api/upload/validate/transactions',
            data={'file':(io.BytesIO(candidate),'older-ledger.xlsx')})
        self.assertEqual(preview.status_code,200,preview.data)
        self.assertTrue(preview.json['can_import'])
        self.assertTrue(preview.json['comparison']['requires_acknowledgement'])
        self.assertEqual(preview.json['comparison']['removed'],1)
        self.assertIn(removed_id,preview.json['comparison']['removed_sample'])
        blocked=self.client.post('/api/upload/transactions',
            data={'file':(io.BytesIO(candidate),'older-ledger.xlsx')})
        self.assertEqual(blocked.status_code,409,blocked.data)
        accepted=self.client.post('/api/upload/transactions',
            data={'allow_regression':'1','file':(io.BytesIO(candidate),'older-ledger.xlsx')})
        self.assertEqual(accepted.status_code,200,accepted.data)
        imported=[row for row in self.client.get('/api/data').json['transactions'] if str(row['id'])==removed_id]
        self.assertEqual(imported,[])
        event=next(row for row in self.client.get('/api/audit-events').json if row['action']=='Imported workbook')
        self.assertTrue(event['details']['regression_acknowledged'])

    def test_reconciliation_exposes_note_level_differences(self):
        reconciliation=self.client.get('/api/data').json['reconciliation']
        self.assertEqual(reconciliation['matched']+reconciliation['review'],len(reconciliation['rows']))
        self.assertTrue(all('ledger_outstanding' in row and 'simulation_outstanding' in row
                            for row in reconciliation['rows']))

    def test_transaction_note_override_is_audited_and_reversible(self):
        before=self.client.get('/api/data').json
        transaction=next(row for row in before['transactions']
                         if row['status']=='SUCCESSFUL' and row['action'] in ('Investment Committed','Principal Payout','Profit Payout'))
        candidates=[row['loan_code'] for row in before['portfolio'] if row['loan_code']!=transaction['note']]
        self.assertTrue(candidates)
        response=self.client.post('/api/transaction-note-override/'+str(transaction['id']),json={'loan_code':candidates[0]})
        self.assertEqual(response.status_code,200,response.data)
        changed=next(row for row in self.client.get('/api/data').json['transactions'] if row['id']==transaction['id'])
        self.assertEqual(changed['note'],candidates[0])
        self.assertEqual(changed['original_note'],transaction['note'])
        events=self.client.get('/api/audit-events').json
        self.assertTrue(any(row['action']=='Assigned note' and row['entity_id']==str(transaction['id']) for row in events))
        self.assertEqual(self.client.delete('/api/transaction-note-override/'+str(transaction['id'])).status_code,200)
        restored=next(row for row in self.client.get('/api/data').json['transactions'] if row['id']==transaction['id'])
        self.assertEqual(restored['note'],transaction['note'])

    def test_expanded_profit_simulation_format_is_accepted(self):
        headers=['No.','Loan Code','Company ID','Issuer Name','Note Name','Product Type',
                 'CTOS/Payment Risk Rating **','Business Description','Investment Amount','Loan Note Size',
                 'Email on Allocation','Email on Disbursement','Investment Exposure',
                 'Total exposure against portfolio <20%','Payment Type **','Term','Tenor Type',
                 'Disbursal Date','Final Repayment Date ','Loan Status','Interest Rate p.a (Gross)',
                 'Interest Rate p.m (Gross)','Interest Rate p.m (Net)','Expected Repayment',
                 'Actual Repayment **','Paid Principal','Unpaid Principal','Gross Profit Earned',
                 'Total Gross Profit','Late Payment Charges','Service Fee ','SST','Net Profit',
                 'Paid Profit','Unpaid Profit','Early Repayment','Early Repayment Date','Remarks','Installment']
        values={
            'No.':1,'Loan Code':'IIF-9998','Company ID':6001,'Issuer Name':'FORMAT TEST SDN BHD',
            'Note Name':'Expanded Profit Test','Product Type':'Islamic Invoice Financing (IIF) - Receivables',
            'Investment Amount':100000,'Loan Note Size':200000,'Investment Exposure':.5,
            'Payment Type **':'Profit Only','Term':3,'Tenor Type':'Months','Disbursal Date':'2026-09-01',
            'Final Repayment Date ':'2026-12-01','Loan Status':'Disbursed','Interest Rate p.a (Gross)':.15,
            'Interest Rate p.m (Gross)':.0125,'Interest Rate p.m (Net)':.01,'Expected Repayment':103000,
            'Actual Repayment **':0,'Paid Principal':0,'Unpaid Principal':100000,
            'Gross Profit Earned':3750,'Total Gross Profit':3800,'Late Payment Charges':50,
            'Service Fee ':760,'SST':20,'Net Profit':3020,'Paid Profit':1000,'Unpaid Profit':2020,
            'Installment':1250
        }
        wb=openpyxl.Workbook();ws=wb.active;ws.title='Query result';ws.append(headers);ws.append([values.get(h) for h in headers])
        schedule=wb.create_sheet('Sheet1')
        for column,value in enumerate(['Notes','Monthly','Gross/Nett','Tenure','Final Repayment Date','Pmt Type','Sept'],1):
            schedule.cell(4,column).value=value
        content=io.BytesIO();wb.save(content);wb.close();content.seek(0)
        response=self.client.post('/api/upload/simulation',data={
            'as_of':'2026-09-15','allow_regression':'1','file':(content,'expanded.xlsx')})
        self.assertEqual(response.status_code,200,response.data)
        row=self.client.get('/api/data').json['portfolio'][0]
        self.assertEqual(row['report_format'],'expanded-profit')
        self.assertEqual(row['gross_profit'],3800)
        self.assertEqual(row['late_profit'],50)
        self.assertEqual(row['sst'],20)

    def test_empty_workspace_accepts_first_import(self):
        with tempfile.TemporaryDirectory() as directory:
            initialize(Path(directory))
            with patch.dict(os.environ,{'MISB_DATA_DIR':directory}):
                response=self.client.get('/api/data')
                self.assertEqual(response.status_code,200)
                self.assertEqual(response.json['summary']['portfolio_count'],0)
                response=self.client.post('/api/upload/simulation',data={'file':(io.BytesIO((self.seed/'simulation.xlsx').read_bytes()),'simulation.xlsx'),'as_of':'2026-08-15'})
                self.assertEqual(response.status_code,200)

    def test_login_and_cross_origin_protection(self):
        with patch.dict(os.environ,{'MISB_LOCAL_PREVIEW':'0','APP_PASSWORD':'test-password-long','SESSION_SECRET':'x'*48}):
            self.assertEqual(self.client.get('/api/data').status_code,401)
            self.assertEqual(self.client.get('/').status_code,302)
            self.assertEqual(self.client.get('/login').status_code,200)
            self.assertEqual(self.client.post('/api/login',json={'password':'bad'}).status_code,401)
            self.assertEqual(self.client.post('/api/login',json={'password':'test-password-long'}).status_code,200)
            self.assertEqual(self.client.get('/api/data').status_code,200)
            self.assertEqual(self.client.post('/api/settings',json={'issuer_limit':100},headers={'Origin':'https://evil.example'}).status_code,403)
            self.assertEqual(self.client.post('/api/logout').status_code,200)
            self.assertEqual(self.client.get('/api/data').status_code,401)

    def test_invalid_json_and_static_privacy(self):
        self.assertEqual(self.client.post('/api/settings',json=[]).status_code,400)
        self.assertEqual(self.client.get('/data/transactions.xlsx').status_code,404)
        response=self.client.get('/')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.headers['Cache-Control'],'no-store')


if __name__=='__main__':
    unittest.main()
